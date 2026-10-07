"""S07-01：人工核定数据合同、读取与校验。

人工核定文件是**评估的唯一标准答案来源**（[plan/04 §1](../../../plan/04-测试与质量评估.md)）：
由用户逐条标注，程序只负责读、校验和记录，**不从 agent 预测反填标签**。

合同（S07 阶段文档 §0.1 第 2 项、§0.3 第 1 条）：

- 主工作表（沿用输入约定 `QA_REF`）：13 列输入原件（编号、q、a、ref1—ref10）
  加 5 列核定列——`人工判断`／`人工理由`／`关键ref`／`是否排除`／`排除理由`；
  第 1 行必须是表头，自上而下即记录顺序。
- `人工判断` 只允许 [schemas.judgement.LABELS] 的三个中文字符串；被排除的行
  **留空**并且 `排除理由` 必填（`是否排除` 写「是」）；既无标签、又未标排除的行
  视为**未核定**，属错误输入。
- 另一张工作表 `核定元数据`：两列「项／值」，四项齐全——`数据版本`、`核定人`、
  `核定日期`（``YYYY-MM-DD``）、`原始文件sha256前12位`（12 位小写十六进制）。

校验一次报出**全部问题**（含行号），不逐条中断——用户要按提示一次改完 100 行，
而不是改一条跑一条。`人工理由` 不是错误项：plan/04 §1 要求保存它，但缺了不妨碍
计分，因此记为 `GoldSet.warnings`。真实核定文件只存在于 ``data/evaluation/``
本地（``data/**`` 已被 `.gitignore` 覆盖，**不进公开仓库**）。
"""

from __future__ import annotations

import re
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from ..excel.reader import file_digest, select_sheet
from ..schemas.judgement import LABELS
from ..schemas.qa import (
    A_COLUMN,
    ID_COLUMN,
    INPUT_COLUMNS,
    Q_COLUMN,
    REF_FIELDS,
    QARecord,
    cell_to_text,
)

#: 核定文件合同版本；改动列名、元数据项或校验规则时递增。
GOLD_CONTRACT_VERSION = "1.0"

#: 主工作表名：与输入读取的约定一致（`excel.reader.PREFERRED_SHEET`）。
GOLD_SHEET = "QA_REF"

GOLD_JUDGEMENT_COLUMN = "人工判断"
GOLD_REASON_COLUMN = "人工理由"
GOLD_KEY_REF_COLUMN = "关键ref"
GOLD_EXCLUDED_COLUMN = "是否排除"
GOLD_EXCLUSION_REASON_COLUMN = "排除理由"

#: 五列核定列的固定顺序：紧跟 13 列输入原件之后。
GOLD_EXTRA_COLUMNS: tuple[str, ...] = (
    GOLD_JUDGEMENT_COLUMN,
    GOLD_REASON_COLUMN,
    GOLD_KEY_REF_COLUMN,
    GOLD_EXCLUDED_COLUMN,
    GOLD_EXCLUSION_REASON_COLUMN,
)

#: 核定文件的完整列顺序（18 列）。
GOLD_COLUMNS: tuple[str, ...] = (*INPUT_COLUMNS, *GOLD_EXTRA_COLUMNS)

#: 元数据工作表名与四个必需项。
GOLD_META_SHEET = "核定元数据"
META_DATA_VERSION = "数据版本"
META_ANNOTATED_BY = "核定人"
META_ANNOTATED_ON = "核定日期"
META_SOURCE_SHA256 = "原始文件sha256前12位"

GOLD_META_KEYS: tuple[str, ...] = (
    META_DATA_VERSION,
    META_ANNOTATED_BY,
    META_ANNOTATED_ON,
    META_SOURCE_SHA256,
)

#: `是否排除` 的两个取值；**留空等同「否」**（见 §0.3 第 1 条的「未核定」判定）。
EXCLUDED_YES = "是"
EXCLUDED_NO = "否"

#: 关键 ref 的多值分隔符：中英文逗号、顿号、分号、空白都接受。
_KEY_REF_SPLIT = re.compile(r"[,，、;；\s]+")

#: 原始文件 sha256 前 12 位：12 位小写十六进制。
_SHA256_PREFIX_RE = re.compile(r"^[0-9a-f]{12}$")

_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")

#: 一次最多打印的问题条数（超出只报总数，避免 100 行全错时刷屏）。
MAX_REPORTED_ISSUES = 30


# --------------------------------------------------------------------------
# 问题与错误
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class GoldIssue:
    """一条核定文件问题；``source_row`` 为工作表中的行号（从 1 数起）。"""

    code: str
    message: str
    source_row: int | None = None

    def render(self) -> str:
        if self.source_row is None:
            return f"[{self.code}] {self.message}"
        return f"第 {self.source_row} 行 [{self.code}] {self.message}"


class GoldError(Exception):
    """核定数据层错误基类。"""


class GoldReadError(GoldError):
    """文件本身读不了（不存在、不是工作簿、工作表无法确定）。"""


class GoldValidationError(GoldError):
    """内容不合合同；一次携带**全部**问题。"""

    def __init__(self, problems: Sequence[GoldIssue]) -> None:
        self.problems: tuple[GoldIssue, ...] = tuple(problems)
        super().__init__(self.render())

    def render(self) -> str:
        head = f"核定文件有 {len(self.problems)} 处问题："
        shown = [issue.render() for issue in self.problems[:MAX_REPORTED_ISSUES]]
        extra = len(self.problems) - len(shown)
        if extra > 0:
            shown.append(f"……另有 {extra} 处问题")
        return "\n".join([head, *shown])


# --------------------------------------------------------------------------
# 结果类型
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class GoldMetadata:
    """核定元数据；四项一一对应 `GOLD_META_KEYS`。"""

    data_version: str
    annotated_by: str
    annotated_on: str
    source_sha256_prefix: str

    def as_dict(self) -> dict[str, str]:
        return {
            META_DATA_VERSION: self.data_version,
            META_ANNOTATED_BY: self.annotated_by,
            META_ANNOTATED_ON: self.annotated_on,
            META_SOURCE_SHA256: self.source_sha256_prefix,
        }


@dataclass(frozen=True)
class GoldEntry:
    """一条核定记录：13 列原件 + 五个核定字段。"""

    record: QARecord
    #: 业务编号；校验阶段已保证非空，因此与 `record.record_id` 分开存一份非可选值，
    #: 免得每处用 `record_id` 都要再判一次 None。
    record_id: str
    label: str | None
    reason: str | None
    key_refs: tuple[str, ...]
    excluded: bool
    exclusion_reason: str | None

    @property
    def source_row(self) -> int:
        return self.record.source_row


@dataclass(frozen=True)
class GoldSet:
    """一次成功读取并校验通过的核定文件。"""

    path: Path
    sha256: str
    sheet_name: str
    contract_version: str
    metadata: GoldMetadata
    entries: tuple[GoldEntry, ...]
    warnings: tuple[GoldIssue, ...] = ()
    skipped_blank_rows: int = 0
    ignored_columns: tuple[str, ...] = ()

    @property
    def included(self) -> tuple[GoldEntry, ...]:
        """参与划分与计分的记录（未排除）。"""
        return tuple(entry for entry in self.entries if not entry.excluded)

    @property
    def excluded(self) -> tuple[GoldEntry, ...]:
        """人工排除的记录；保留编号与理由，不参与计分。"""
        return tuple(entry for entry in self.entries if entry.excluded)

    @property
    def record_ids(self) -> tuple[str, ...]:
        return tuple(entry.record_id for entry in self.entries)

    @property
    def included_ids(self) -> tuple[str, ...]:
        return tuple(entry.record_id for entry in self.included)

    @property
    def label_counts(self) -> dict[str, int]:
        """未排除记录的逐标签计数；三个标签都会出现，计数可以为 0。"""
        counts = {label: 0 for label in LABELS}
        for entry in self.included:
            if entry.label is None:  # pragma: no cover - 校验已拦住
                continue
            counts[entry.label] += 1
        return counts

    @property
    def sha256_prefix(self) -> str:
        return self.sha256[:12]


# --------------------------------------------------------------------------
# 读取与校验
# --------------------------------------------------------------------------


def source_sha256_prefix(digest: str) -> str:
    """取 sha256 前 12 位小写，用于元数据比对（批次上传摘要、gold 摘要同法）。"""
    return digest.strip().lower()[:12]


def compare_ids(
    expected: Iterable[str], actual: Iterable[str]
) -> tuple[tuple[str, ...], tuple[str, ...]]:
    """返回 ``(缺失, 多出)`` 两个按字典序排列的编号差集。

    用于「批次编号集合 vs 划分清单编号集合」这类**同构**校验：两个差集都为空
    才算一致。缺失=清单里有而批次没有，多出=批次里有而清单没有。
    """
    expected_set = set(expected)
    actual_set = set(actual)
    missing = tuple(sorted(expected_set - actual_set))
    extra = tuple(sorted(actual_set - expected_set))
    return missing, extra


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _sheet_rows(workbook: object, name: str) -> list[tuple[object, ...]]:
    return list(workbook[name].iter_rows(values_only=True))  # type: ignore[index]


def _cell(row: Sequence[object], position: int | None) -> object:
    if position is None or position >= len(row):
        return None
    return row[position]


def _header_positions(header: Sequence[object]) -> dict[str, int | None]:
    text = [cell_to_text(value) for value in header]
    return {column: (text.index(column) if column in text else None) for column in GOLD_COLUMNS}


def _read_metadata(workbook: object, problems: list[GoldIssue]) -> GoldMetadata | None:
    names = [sheet.title for sheet in workbook.worksheets]  # type: ignore[attr-defined]
    if GOLD_META_SHEET not in names:
        problems.append(
            GoldIssue(
                "GOLD_META_MISSING",
                f"找不到元数据工作表 {GOLD_META_SHEET}；需要四项："
                f"{'、'.join(GOLD_META_KEYS)}",
            )
        )
        return None

    values: dict[str, str] = {}
    for row in _sheet_rows(workbook, GOLD_META_SHEET):
        key = cell_to_text(_cell(row, 0))
        if key is None or key not in GOLD_META_KEYS:
            continue  # 表头行与无关行都跳过
        text = cell_to_text(_cell(row, 1))
        if text is not None:
            values[key] = text

    missing = [key for key in GOLD_META_KEYS if not values.get(key)]
    if missing:
        problems.append(
            GoldIssue("GOLD_META_INCOMPLETE", f"元数据缺少或为空：{'、'.join(missing)}")
        )

    annotated_on = values.get(META_ANNOTATED_ON, "")
    if annotated_on and not _DATE_RE.match(annotated_on):
        problems.append(
            GoldIssue(
                "GOLD_META_INVALID_DATE",
                f"{META_ANNOTATED_ON} 应为 YYYY-MM-DD，实际为 {annotated_on!r}",
            )
        )

    prefix = values.get(META_SOURCE_SHA256, "")
    if prefix and not _SHA256_PREFIX_RE.match(prefix.strip().lower()):
        problems.append(
            GoldIssue(
                "GOLD_META_INVALID_SHA256",
                f"{META_SOURCE_SHA256} 应为 12 位小写十六进制，实际为 {prefix!r}",
            )
        )

    if problems:
        return None
    return GoldMetadata(
        data_version=values[META_DATA_VERSION],
        annotated_by=values[META_ANNOTATED_BY],
        annotated_on=values[META_ANNOTATED_ON],
        source_sha256_prefix=values[META_SOURCE_SHA256].strip().lower(),
    )


def _key_refs(text: str | None) -> tuple[str, ...]:
    if text is None:
        return ()
    return tuple(token for token in _KEY_REF_SPLIT.split(text) if token)


def _parse_entry(
    raw: Sequence[object],
    positions: Mapping[str, int | None],
    *,
    source_row: int,
    order_index: int,
    problems: list[GoldIssue],
    warnings: list[GoldIssue],
) -> GoldEntry | None:
    """校验一行并构造条目；有致命问题时返回 ``None``（问题已记入 ``problems``）。"""
    cells = {column: _cell(raw, positions.get(column)) for column in GOLD_COLUMNS}
    record_id = cell_to_text(cells[ID_COLUMN])
    judgement = cell_to_text(cells[GOLD_JUDGEMENT_COLUMN])
    reason = cell_to_text(cells[GOLD_REASON_COLUMN])
    key_ref_text = cell_to_text(cells[GOLD_KEY_REF_COLUMN])
    excluded_text = cell_to_text(cells[GOLD_EXCLUDED_COLUMN])
    exclusion_reason = cell_to_text(cells[GOLD_EXCLUSION_REASON_COLUMN])

    fatal = False

    if record_id is None:
        problems.append(GoldIssue("GOLD_ID_MISSING", f"{ID_COLUMN} 不能为空", source_row))
        fatal = True

    if excluded_text not in (None, EXCLUDED_YES, EXCLUDED_NO):
        problems.append(
            GoldIssue(
                "GOLD_EXCLUDED_INVALID",
                f"{GOLD_EXCLUDED_COLUMN} 只接受 {EXCLUDED_YES!r}／{EXCLUDED_NO!r} 或留空，"
                f"实际为 {excluded_text!r}",
                source_row,
            )
        )
        fatal = True

    excluded = excluded_text == EXCLUDED_YES

    if excluded:
        if judgement is not None:
            problems.append(
                GoldIssue(
                    "GOLD_EXCLUDED_HAS_LABEL",
                    f"被排除的行 {GOLD_JUDGEMENT_COLUMN} 必须留空（该行不计分）",
                    source_row,
                )
            )
            fatal = True
        if exclusion_reason is None:
            problems.append(
                GoldIssue(
                    "GOLD_EXCLUSION_REASON_MISSING",
                    f"标为{EXCLUDED_YES}的行必须填写 {GOLD_EXCLUSION_REASON_COLUMN}",
                    source_row,
                )
            )
            fatal = True
    elif judgement is None:
        problems.append(
            GoldIssue(
                "GOLD_UNDECIDED",
                f"既没有 {GOLD_JUDGEMENT_COLUMN}，也没有标 {GOLD_EXCLUDED_COLUMN}={EXCLUDED_YES}；"
                "未核定的行不能作为标准答案",
                source_row,
            )
        )
        fatal = True
    elif judgement not in LABELS:
        problems.append(
            GoldIssue(
                "GOLD_LABEL_INVALID",
                f"{GOLD_JUDGEMENT_COLUMN} 必须是 {'／'.join(LABELS)} 之一，实际为 {judgement!r}",
                source_row,
            )
        )
        fatal = True
    elif reason is None:
        warnings.append(
            GoldIssue(
                "GOLD_REASON_MISSING",
                f"未填写 {GOLD_REASON_COLUMN}（plan/04 §1 要求保存人工理由，但不影响计分）",
                source_row,
            )
        )

    key_refs = _key_refs(key_ref_text)
    unknown_refs = [ref for ref in key_refs if ref not in REF_FIELDS]
    if unknown_refs:
        problems.append(
            GoldIssue(
                "GOLD_KEY_REF_UNKNOWN",
                f"{GOLD_KEY_REF_COLUMN} 只能是 {REF_FIELDS[0]}—{REF_FIELDS[-1]}，"
                f"无法识别：{'、'.join(unknown_refs)}",
                source_row,
            )
        )
        fatal = True

    refs: dict[str, str | None] = {}
    for field in REF_FIELDS:
        refs[field] = cell_to_text(cells[field])

    empty_key_refs = [ref for ref in key_refs if refs.get(ref) is None]
    if empty_key_refs:
        problems.append(
            GoldIssue(
                "GOLD_KEY_REF_EMPTY",
                f"{GOLD_KEY_REF_COLUMN} 指向的资料为空：{'、'.join(empty_key_refs)}",
                source_row,
            )
        )
        fatal = True

    q = cell_to_text(cells[Q_COLUMN])
    a = cell_to_text(cells[A_COLUMN])
    missing_fields = [
        column for column, value in ((Q_COLUMN, q), (A_COLUMN, a)) if value is None
    ]
    if missing_fields:
        problems.append(
            GoldIssue(
                "GOLD_INPUT_FIELD_MISSING",
                f"核定文件要保留 13 列输入原件，缺少：{'、'.join(missing_fields)}",
                source_row,
            )
        )
        fatal = True

    if fatal:
        return None

    return GoldEntry(
        record=QARecord(
            record_id=record_id,
            source_row=source_row,
            order_index=order_index,
            q=q,
            a=a,
            refs=refs,
        ),
        record_id=record_id,
        label=None if excluded else judgement,
        reason=reason,
        key_refs=key_refs,
        excluded=excluded,
        exclusion_reason=exclusion_reason,
    )


def read_gold(path: str | Path, sheet_name: str | None = None) -> GoldSet:
    """读取并校验人工核定文件。

    成功返回 `GoldSet`；内容不合合同抛 `GoldValidationError`（携带全部问题），
    文件读不了抛 `GoldReadError`。公式单元格按**缓存值**读取（``data_only=True``），
    没有缓存值的公式与空单元格同样视为「未填写」。
    """
    source = Path(path)
    if not source.is_file():
        raise GoldReadError(f"找不到人工核定文件：{source}")

    digest = file_digest(source)
    try:
        workbook = load_workbook(filename=source, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 对损坏文件抛出多种异常
        raise GoldReadError(f"核定文件无法读取：{exc}") from exc

    problems: list[GoldIssue] = []
    warnings: list[GoldIssue] = []
    try:
        sheets = tuple(
            (sheet.title, sheet.sheet_state == "visible")
            for sheet in workbook.worksheets  # type: ignore[attr-defined]
        )
        metadata = _read_metadata(workbook, problems)

        candidates = tuple(item for item in sheets if item[0] != GOLD_META_SHEET)
        selection = select_sheet(candidates, sheet_name)
        if selection.name is None:
            raise GoldReadError(
                f"无法确定核定工作表：{selection.reason}（元数据表 {GOLD_META_SHEET} 已排除在候选之外）"
            )

        rows = _sheet_rows(workbook, selection.name)
        if not rows:
            raise GoldReadError(f"工作表 {selection.name} 是空的，没有表头")

        positions = _header_positions(rows[0])
        missing_columns = [column for column in GOLD_COLUMNS if positions[column] is None]
        if missing_columns:
            problems.append(
                GoldIssue(
                    "GOLD_MISSING_COLUMNS",
                    f"表头缺少列：{'、'.join(missing_columns)}（需 13 列输入原件 + "
                    f"{len(GOLD_EXTRA_COLUMNS)} 列核定列）",
                )
            )
        ignored = tuple(
            text
            for text in (cell_to_text(value) for value in rows[0])
            if text is not None and text not in GOLD_COLUMNS
        )

        entries: list[GoldEntry] = []
        seen: dict[str, int] = {}
        skipped = 0
        data_rows = 0
        if not missing_columns:
            for row_number, raw in enumerate(rows[1:], start=2):
                if all(_blank(value) for value in raw):
                    skipped += 1
                    continue
                data_rows += 1
                entry = _parse_entry(
                    raw,
                    positions,
                    source_row=row_number,
                    order_index=len(entries),
                    problems=problems,
                    warnings=warnings,
                )
                if entry is None:
                    continue
                if entry.record_id in seen:
                    problems.append(
                        GoldIssue(
                            "GOLD_ID_DUPLICATE",
                            f"{ID_COLUMN} 重复：{entry.record_id}（首次出现在第 "
                            f"{seen[entry.record_id]} 行）",
                            row_number,
                        )
                    )
                    continue
                seen[entry.record_id] = row_number
                entries.append(entry)

            # 只在**一行都没有**时报「没有记录」：整表都是坏行时真正的信息是那些
            # 行级问题，再补一条「没有记录」只会把提示语引向错误的方向。
            if data_rows == 0:
                problems.append(
                    GoldIssue("GOLD_NO_ROWS", f"工作表 {selection.name} 里没有任何核定记录")
                )
    finally:
        workbook.close()

    if problems:
        raise GoldValidationError(problems)
    if metadata is None:  # pragma: no cover - problems 非空时已抛出
        raise GoldValidationError(
            [GoldIssue("GOLD_META_MISSING", f"找不到元数据工作表 {GOLD_META_SHEET}")]
        )

    return GoldSet(
        path=source,
        sha256=digest,
        sheet_name=selection.name,
        contract_version=GOLD_CONTRACT_VERSION,
        metadata=metadata,
        entries=tuple(entries),
        warnings=tuple(warnings),
        skipped_blank_rows=skipped,
        ignored_columns=ignored,
    )


# --------------------------------------------------------------------------
# 写：测试与样例用
# --------------------------------------------------------------------------


def build_gold_workbook(
    rows: Iterable[Mapping[str, object]], metadata: Mapping[str, object]
) -> bytes:
    """按合同生成一个核定工作簿（内存字节）。

    供**测试与合成样例**使用：真实核定文件由用户产出，不由程序生成。
    行按列名给出；未给出的列为空。元数据项按 `GOLD_META_KEYS` 顺序写出。

    与 `excel.samples` 一样**不承诺字节稳定**：xlsx 内含时间戳，同一份内容两次
    生成的 sha256 不同。
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = GOLD_SHEET
    sheet.append(list(GOLD_COLUMNS))
    for row in rows:
        sheet.append([row.get(column) for column in GOLD_COLUMNS])

    meta = workbook.create_sheet(title=GOLD_META_SHEET)
    meta.append(["项", "值"])
    for key in GOLD_META_KEYS:
        meta.append([key, metadata.get(key)])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def write_gold_workbook(
    path: str | Path, rows: Iterable[Mapping[str, object]], metadata: Mapping[str, object]
) -> Path:
    """把 `build_gold_workbook` 的字节写到磁盘并返回路径。"""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(build_gold_workbook(rows, metadata))
    return target


__all__ = [
    "EXCLUDED_NO",
    "EXCLUDED_YES",
    "GOLD_COLUMNS",
    "GOLD_CONTRACT_VERSION",
    "GOLD_EXCLUDED_COLUMN",
    "GOLD_EXCLUSION_REASON_COLUMN",
    "GOLD_EXTRA_COLUMNS",
    "GOLD_JUDGEMENT_COLUMN",
    "GOLD_KEY_REF_COLUMN",
    "GOLD_META_KEYS",
    "GOLD_META_SHEET",
    "GOLD_REASON_COLUMN",
    "GOLD_SHEET",
    "META_ANNOTATED_BY",
    "META_ANNOTATED_ON",
    "META_DATA_VERSION",
    "META_SOURCE_SHA256",
    "GoldEntry",
    "GoldError",
    "GoldIssue",
    "GoldMetadata",
    "GoldReadError",
    "GoldSet",
    "GoldValidationError",
    "build_gold_workbook",
    "compare_ids",
    "read_gold",
    "source_sha256_prefix",
    "write_gold_workbook",
]
