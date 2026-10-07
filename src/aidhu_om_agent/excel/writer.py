"""S05-02：导出工作簿的写出（四表 Excel）。

本模块只做一件事：把**已经编码好的字符串**写成 xlsx。快照装载、行内容与
截断/防护判断都不在这里（见 `services/exports.py` 与 `schemas/export.py`）。

两条硬规则：

- **单元格里不允许出现公式**。调用方应当先用
  `schemas.export.excel_cell_text` 编码；这里再加一道断言，把一个漏编码的
  字符串当场变成失败，而不是生成一份打开就求值的文件。
- **不扫全表**。列宽只看前若干行——真正长的单元格文本可到 32767 字符，
  为了排版扫完整批没有意义。
"""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..schemas.export import FORMULA_TRIGGER_PREFIXES

#: 允许写入的单元格值；`序号` 写数字便于在 Excel 里排序，其余都是文本。
CellValue = str | int | float | None


class ExcelWriteError(RuntimeError):
    """写入导出工作簿时的可报告错误。"""


@dataclass(frozen=True)
class SheetData:
    """一张待写出的工作表：名字、表头列序、数据行。"""

    name: str
    columns: tuple[str, ...]
    rows: tuple[tuple[CellValue, ...], ...] = ()


#: 列宽上限与下限（按字符宽度计）；下限保证表头不被挤成 ``###``。
MAX_COLUMN_WIDTH = 60
MIN_COLUMN_WIDTH = 10

#: 估算列宽时最多看多少行；超出的行不参与排版。
WIDTH_SAMPLE_ROWS = 200


def _display_width(value: CellValue) -> int:
    """近似显示宽度；CJK 等全角字符按两列计。只用于列宽，不参与任何判定。"""
    if value is None:
        return 0
    text = str(value)
    return sum(2 if ord(char) >= 0x1100 else 1 for char in text)


def _column_widths(sheet: SheetData) -> list[int]:
    sample = sheet.rows[:WIDTH_SAMPLE_ROWS]
    widths: list[int] = []
    for index, column in enumerate(sheet.columns):
        widest = _display_width(column)
        for row in sample:
            if index < len(row):
                widest = max(widest, _display_width(row[index]))
        widths.append(min(max(widest + 2, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH))
    return widths


def _write_cell(cell, value: CellValue, *, sheet: str, row: int) -> None:
    """写一个单元格；**文本永远当文本**，出现公式即报错。"""
    if isinstance(value, str) and value.startswith(FORMULA_TRIGGER_PREFIXES):
        raise ExcelWriteError(
            f"{sheet} 第 {row} 行的单元格以 {value[:1]!r} 开头且未经文本前缀防护；"
            "导出不允许写出公式单元格"
        )
    cell.value = value
    # 双保险：即使将来 openpyxl 改了判定规则，也不让公式落到文件里。
    if cell.data_type == "f":
        raise ExcelWriteError(
            f"{sheet} 第 {row} 行的单元格被当作公式写入（值以 {str(value)[:1]!r} 开头）"
        )


def build_workbook(sheets: Sequence[SheetData]) -> Workbook:
    """按给定顺序装配工作簿；表头加粗、冻结首行、按内容估算列宽。"""
    if not sheets:
        raise ExcelWriteError("导出工作簿至少要有一张工作表")
    names = [sheet.name for sheet in sheets]
    if len(set(names)) != len(names):
        raise ExcelWriteError(f"工作表名重复：{names}")

    workbook = Workbook()
    # 默认建出的空表会变成一个多余的标签页，直接换成第一张真实表。
    default = workbook.active
    workbook.remove(default)

    header_font = Font(bold=True)
    for sheet in sheets:
        worksheet = workbook.create_sheet(title=sheet.name)
        for index, column in enumerate(sheet.columns, start=1):
            cell = worksheet.cell(row=1, column=index)
            _write_cell(cell, column, sheet=sheet.name, row=1)
            cell.font = header_font
        for offset, row in enumerate(sheet.rows, start=2):
            if len(row) != len(sheet.columns):
                raise ExcelWriteError(
                    f"{sheet.name} 第 {offset} 行有 {len(row)} 列，"
                    f"表头是 {len(sheet.columns)} 列"
                )
            for index, value in enumerate(row, start=1):
                _write_cell(
                    worksheet.cell(row=offset, column=index),
                    value,
                    sheet=sheet.name,
                    row=offset,
                )
        worksheet.freeze_panes = "A2"
        for index, width in enumerate(_column_widths(sheet), start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = width

    return workbook


def write_workbook(path: Path, sheets: Iterable[SheetData]) -> Path:
    """写出工作簿到 ``path``；父目录不存在时创建。

    写的是**临时文件**还是最终路径由调用方决定：导出采用「两份都写好再改名」的
    发布顺序（plan/09 §8），所以这里不负责改名，也不负责登记。
    """
    workbook = build_workbook(tuple(sheets))
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return target


__all__ = [
    "CellValue",
    "ExcelWriteError",
    "MAX_COLUMN_WIDTH",
    "MIN_COLUMN_WIDTH",
    "SheetData",
    "build_workbook",
    "write_workbook",
]
