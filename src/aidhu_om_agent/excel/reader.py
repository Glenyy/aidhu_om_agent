"""S02-01—S02-03、S02-05、S03-07：13 列 Excel 输入的读取、工作表选择与批次预检。

本模块只做解析与预检：**不调用模型、不写数据库、不改写原文件**。
批次级问题产出 `blockers`（status 为 blocked，不产出可处理记录）；
单条问题产出 `row_errors`，其余有效记录照常返回。

单元格读取约定见 S02 阶段文档 §0：公式按 ``data_only=True`` 取缓存值，
无缓存值或结果为 Excel 错误值的单元格记为不可读，不当作资料文本。

S06-02 起加**上传规模防护**（[plan/08 §3]）：解压总量、可见工作表数在打开工作簿
前后各查一次，选中表的单元格数在物化行之前查。三项都不裁剪资料，超限抛
`WorkbookTooLargeError`（`InputReadError` 的子类，接口层统一映射为 422 `BAD_WORKBOOK`）。
"""

from __future__ import annotations

import hashlib
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from ..config import LimitsConfig, default_limits
from ..schemas.qa import (
    A_COLUMN,
    ID_COLUMN,
    INPUT_COLUMNS,
    INPUT_CONTRACT_VERSION,
    Q_COLUMN,
    REF_FIELDS,
    Blocker,
    ParsedInput,
    PrecheckCounts,
    PrecheckReport,
    PrecheckWarning,
    QARecord,
    RowError,
    cell_to_text,
    is_error_cell,
    is_formula_cell,
    is_numeric_cell,
)

#: 约定优先使用的工作表名。
PREFERRED_SHEET = "QA_REF"

BLOCKER_FILE_TOO_LARGE = "FILE_TOO_LARGE"
BLOCKER_SHEET_AMBIGUOUS = "SHEET_AMBIGUOUS"
BLOCKER_MISSING_COLUMNS = "MISSING_COLUMNS"
BLOCKER_DUPLICATE_COLUMNS = "DUPLICATE_COLUMNS"
BLOCKER_DUPLICATE_RECORD_ID = "DUPLICATE_RECORD_ID"
BLOCKER_TOO_MANY_RECORDS = "TOO_MANY_RECORDS"
BLOCKER_NO_VALID_RECORDS = "NO_VALID_RECORDS"

WARNING_EXTRA_COLUMNS = "EXTRA_COLUMNS_UNUSED"
WARNING_NUMERIC_RECORD_ID = "NUMERIC_RECORD_ID_LEADING_ZEROS"

#: 必须能从单元格读出文本的列；ref 为空合法，不在此列。
REQUIRED_TEXT_COLUMNS: tuple[str, ...] = (ID_COLUMN, Q_COLUMN, A_COLUMN)

#: 解压总量上限的倍数：实际上限 = 该倍数 × ``limits.max_upload_bytes``
#: （S06 阶段文档 §0.3 第 7 项）。跟着配置联动，不写死 400 MiB。
MAX_EXPANDED_RATIO = 8

#: 单个工作表允许读取的单元格数上限（行 × 列）。
MAX_SHEET_CELLS = 500_000

#: 允许的**可见**工作表数上限。隐藏表不计入，但它们的内容同样会被读取，
#: 所以隐藏表不设上限这一条由「解压总量」兜底。
MAX_VISIBLE_SHEETS = 50


class InputReadError(Exception):
    """输入文件不存在或工作簿无法读取；属于调用错误，不是预检结论。"""


class WorkbookTooLargeError(InputReadError):
    """规模防护命中：解压总量、单元格数或可见工作表数超限。

    与「文件损坏」归为同一类**调用错误**（HTTP 422 `BAD_WORKBOOK`），不是预检结论：
    预检报告描述的是内容是否符合输入合同，而规模超限在打开工作簿之前/读取单元格
    之前就被拒绝，根本没有产出报告的机会。超限一律**不裁剪资料**——少读一部分再
    给出结论，等于用不完整输入下判断。
    """


@dataclass(frozen=True)
class SheetSelection:
    """工作表选择结果；``name`` 为 None 表示无法确定。"""

    name: str | None
    reason: str
    candidates: tuple[tuple[str, bool], ...] = ()


# --------------------------------------------------------------------------
# 工具
# --------------------------------------------------------------------------


def file_digest(path: Path) -> str:
    """输入文件的 sha256；用于固定预检对应的文件内容。"""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def select_sheet(
    sheets: Iterable[tuple[str, bool]], requested: str | None = None
) -> SheetSelection:
    """按显式指定 → QA_REF → 唯一可见表的顺序选择工作表。

    ``sheets`` 为 (表名, 是否可见) 序列。无法确定时返回 ``name=None``，
    并带上全部候选表，供人工明确指定。
    """
    candidates = tuple(sheets)
    names = tuple(name for name, _ in candidates)

    if requested is not None:
        if requested in names:
            return SheetSelection(requested, f"按显式指定的工作表 {requested!r}", candidates)
        return SheetSelection(
            None,
            f"工作簿中没有工作表 {requested!r}；实际工作表：{_name_list(names)}",
            candidates,
        )

    if PREFERRED_SHEET in names:
        return SheetSelection(
            PREFERRED_SHEET, f"按约定优先使用 {PREFERRED_SHEET}", candidates
        )

    visible = tuple(name for name, visible in candidates if visible)
    if len(visible) == 1:
        return SheetSelection(visible[0], f"没有 {PREFERRED_SHEET}，仅一个可见工作表", candidates)

    return SheetSelection(
        None,
        f"没有 {PREFERRED_SHEET}，且有 {len(visible)} 个可见工作表，无法自动确定；"
        f"可见工作表：{_name_list(visible)}",
        candidates,
    )


def _name_list(names: Sequence[str]) -> str:
    return "、".join(names) if names else "（无）"


# --------------------------------------------------------------------------
# S06-02：上传规模防护（[plan/08 §3] 的「解压总量、单元格及工作表读取保护」）
# --------------------------------------------------------------------------


def expanded_size_bytes(path: str | Path) -> int:
    """工作簿解压后的总字节数；**只读 zip 中央目录，不解压任何成员**。

    取的是各成员声明的 ``file_size``。xlsx 是 zip，一个几十 KB 的文件可以声明解压
    后几十 GB；这一步在 `openpyxl.load_workbook` 之前把这种文件挡掉，代价只有一次
    目录读取。

    不是 zip（也就是不是 xlsx）时在这里就报「工作簿无法读取」——与 openpyxl 自己的
    失败落在同一条路径上，调用方不需要多一层分支。
    """
    try:
        with zipfile.ZipFile(Path(path)) as archive:
            return sum(info.file_size for info in archive.infolist())
    except (zipfile.BadZipFile, OSError) as exc:
        raise InputReadError(f"工作簿无法读取：{exc}") from exc


def guard_expanded_size(path: str | Path, limits: LimitsConfig) -> int:
    """解压总量防护；返回实际解压字节数，超限抛 `WorkbookTooLargeError`。"""
    actual = expanded_size_bytes(path)
    budget = MAX_EXPANDED_RATIO * limits.max_upload_bytes
    if actual > budget:
        raise WorkbookTooLargeError(
            f"工作簿解压后合计 {actual} 字节，超过上限 {budget} 字节"
            f"（{MAX_EXPANDED_RATIO} × 上传上限 {limits.max_upload_bytes} 字节）；"
            "不裁剪资料"
        )
    return actual


def guard_visible_sheets(sheets: Sequence[tuple[str, bool]]) -> None:
    """可见工作表数防护；超限抛 `WorkbookTooLargeError`。

    隐藏表不计入：界面上只让用户从可见表里选，隐藏表通常是被刻意留着的旧数据。
    """
    visible = sum(1 for _, is_visible in sheets if is_visible)
    if visible > MAX_VISIBLE_SHEETS:
        raise WorkbookTooLargeError(
            f"工作簿有 {visible} 个可见工作表，超过上限 {MAX_VISIBLE_SHEETS} 个；"
            "不裁剪内容"
        )


def guard_sheet_cells(workbook: object, name: str) -> int:
    """**选中**工作表的单元格数防护；返回行 × 列，超限抛 `WorkbookTooLargeError`。

    只查被读取的那张表：文件里另有一张用不到的大表，不该阻断当前这次预检。
    检查发生在 `_sheet_rows` 真的物化全部行**之前**，所以它挡住的正是内存占用。
    """
    sheet = workbook[name]  # type: ignore[index]
    cells = (sheet.max_row or 0) * (sheet.max_column or 0)
    if cells > MAX_SHEET_CELLS:
        raise WorkbookTooLargeError(
            f"工作表 {name!r} 有约 {cells} 个单元格，超过上限 {MAX_SHEET_CELLS} 个；"
            "不裁剪资料"
        )
    return cells


# --------------------------------------------------------------------------
# S03-07：上传接口使用的工作表清单
# --------------------------------------------------------------------------


def sheet_catalog(
    path: str | Path, *, limits: LimitsConfig | None = None
) -> tuple[tuple[str, bool], ...]:
    """列出工作簿中的全部工作表（表名 + 是否可见）。

    S03-07 的上传接口需要在尚未预检时展示可选工作表；本函数与 `precheck`
    共用同一份 `_sheet_list`，不重复实现工作簿读取，也不产出预检结论。

    S06-02 起同时承担**上传即生效**的两项规模防护（解压总量、可见工作表数）：
    上传时就拒绝，比等到预检才拒绝少留一份无用文件。``limits`` 省略时使用内置
    保护值，与 `precheck` 同一约定；调用方应传 `config.limits`。
    """
    effective_limits = limits or default_limits()
    source = Path(path)
    if not source.is_file():
        raise InputReadError(f"找不到输入文件：{source}")

    guard_expanded_size(source, effective_limits)
    try:
        workbook = load_workbook(filename=source, read_only=True, data_only=False)
    except Exception as exc:  # openpyxl 对损坏文件抛出多种异常
        raise InputReadError(f"工作簿无法读取：{exc}") from exc
    try:
        sheets = _sheet_list(workbook)
    finally:
        workbook.close()

    guard_visible_sheets(sheets)
    return sheets


def _at(row: Sequence[object], position: int) -> object:
    return row[position] if position < len(row) else None


def _is_blank_cell(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _sheet_rows(workbook: object, name: str) -> list[tuple[object, ...]]:
    return list(workbook[name].iter_rows(values_only=True))  # type: ignore[index]


def _sheet_list(workbook: object) -> tuple[tuple[str, bool], ...]:
    return tuple(
        (sheet.title, sheet.sheet_state == "visible")
        for sheet in workbook.worksheets  # type: ignore[attr-defined]
    )


# --------------------------------------------------------------------------
# 预检
# --------------------------------------------------------------------------


def precheck(
    path: str | Path,
    sheet_name: str | None = None,
    *,
    limits: LimitsConfig | None = None,
) -> ParsedInput:
    """读取工作簿并产出解析快照。

    ``limits`` 省略时使用内置保护值；调用方（如 CLI）应优先传入配置中的
    ``config.limits``，本函数不自行读取配置文件。

    规模防护分三层：压缩字节数超限是**预检结论**（`blocked` + `FILE_TOO_LARGE`，
    维持 S02 的既有行为）；解压总量与可见工作表数超限、以及选中表的单元格数超限
    都是**调用错误**，抛 `WorkbookTooLargeError`（HTTP 422 `BAD_WORKBOOK`）。
    """
    effective_limits = limits or default_limits()
    source = Path(path)
    if not source.is_file():
        raise InputReadError(f"找不到输入文件：{source}")

    size_bytes = source.stat().st_size
    digest = file_digest(source)

    if size_bytes > effective_limits.max_upload_bytes:
        return _blocked(
            sheet_name=None,
            digest=digest,
            size_bytes=size_bytes,
            counts=PrecheckCounts(),
            blockers=(
                Blocker(
                    code=BLOCKER_FILE_TOO_LARGE,
                    message=(
                        f"文件大小 {size_bytes} 字节超过上限 "
                        f"{effective_limits.max_upload_bytes} 字节；不截取内容"
                    ),
                    field="file_size_bytes",
                ),
            ),
        )

    guard_expanded_size(source, effective_limits)

    try:
        workbook = load_workbook(filename=source, read_only=True, data_only=False)
    except Exception as exc:  # openpyxl 对损坏文件抛出多种异常
        raise InputReadError(f"工作簿无法读取：{exc}") from exc

    try:
        sheets = _sheet_list(workbook)
        guard_visible_sheets(sheets)
        selection = select_sheet(sheets, sheet_name)
        if selection.name is None:
            return _blocked(
                sheet_name=None,
                digest=digest,
                size_bytes=size_bytes,
                counts=PrecheckCounts(),
                blockers=(
                    Blocker(code=BLOCKER_SHEET_AMBIGUOUS, message=selection.reason, field="sheet_name"),
                ),
            )
        guard_sheet_cells(workbook, selection.name)
        raw_rows = _sheet_rows(workbook, selection.name)
    finally:
        workbook.close()

    header_text = [cell_to_text(cell) or "" for cell in (raw_rows[0] if raw_rows else ())]

    index_of: dict[str, int] = {}
    duplicated: list[str] = []
    for position, name in enumerate(header_text):
        if not name:
            continue
        if name in index_of:
            duplicated.append(name)
        else:
            index_of[name] = position

    missing = [column for column in INPUT_COLUMNS if column not in index_of]
    if missing:
        return _blocked(
            sheet_name=selection.name,
            digest=digest,
            size_bytes=size_bytes,
            counts=PrecheckCounts(),
            blockers=(
                Blocker(
                    code=BLOCKER_MISSING_COLUMNS,
                    message=f"第 1 行缺少必要列：{'、'.join(missing)}",
                    field=",".join(missing),
                ),
            ),
        )

    if duplicated:
        names = "、".join(sorted(set(duplicated)))
        return _blocked(
            sheet_name=selection.name,
            digest=digest,
            size_bytes=size_bytes,
            counts=PrecheckCounts(),
            blockers=(
                Blocker(
                    code=BLOCKER_DUPLICATE_COLUMNS,
                    message=f"第 1 行存在重名列，列映射有歧义：{names}",
                    field=names,
                ),
            ),
        )

    data_rows = raw_rows[1:]
    cached_rows = _cached_data_rows(source, selection.name, data_rows, index_of)

    records: list[QARecord] = []
    row_errors: list[RowError] = []
    blank_rows = 0
    numeric_id_rows: list[int] = []

    for offset, raw_row in enumerate(data_rows):
        source_row = offset + 2  # 第 1 行为表头
        if all(_is_blank_cell(_at(raw_row, index_of[column])) for column in INPUT_COLUMNS):
            blank_rows += 1
            continue

        cached_row = cached_rows[offset] if cached_rows is not None else None
        values, unreadable = _read_row(raw_row, cached_row, index_of)

        if is_numeric_cell(_at(raw_row, index_of[ID_COLUMN])):
            numeric_id_rows.append(source_row)

        order_index = len(records)
        texts = {column: cell_to_text(value) for column, value in values.items()}
        record = QARecord(
            record_id=texts[ID_COLUMN],
            source_row=source_row,
            order_index=order_index,
            q=texts[Q_COLUMN],
            a=texts[A_COLUMN],
            refs={field: texts[field] for field in REF_FIELDS},
        )
        records.append(record)

        reasons = _input_error_reasons(record, unreadable)
        if reasons:
            row_errors.append(
                RowError(
                    source_row=source_row,
                    record_id=record.record_id,
                    reason="；".join(reasons),
                )
            )

    invalid_rows = {error.source_row for error in row_errors}
    valid = len(records) - len(invalid_rows)
    counts = PrecheckCounts(
        total=len(records),
        valid=valid,
        input_invalid=len(invalid_rows),
        skipped_blank_rows=blank_rows,
    )

    blockers = _batch_blockers(records, counts, effective_limits)
    warnings = _warnings(header_text, numeric_id_rows)

    report = PrecheckReport(
        input_contract_version=INPUT_CONTRACT_VERSION,
        status="blocked" if blockers else "passed",
        sheet_name=selection.name,
        file_sha256=digest,
        file_size_bytes=size_bytes,
        counts=counts,
        blockers=blockers,
        row_errors=tuple(row_errors),
        warnings=warnings,
    )
    # blocked 时不产出可处理记录：批次级问题必须先修正。
    return ParsedInput(report=report, records=() if blockers else tuple(records))


def _cached_data_rows(
    source: Path,
    sheet_name: str,
    data_rows: Sequence[Sequence[object]],
    index_of: dict[str, int],
) -> list[Sequence[object]] | None:
    """仅当 13 列中存在公式单元格时，第二次读取以取公式缓存值。"""
    positions = tuple(index_of[column] for column in INPUT_COLUMNS)
    needs_cache = any(
        is_formula_cell(_at(row, position)) for row in data_rows for position in positions
    )
    if not needs_cache:
        return None

    try:
        workbook = load_workbook(filename=source, read_only=True, data_only=True)
    except Exception as exc:
        raise InputReadError(f"工作簿无法以缓存值读取：{exc}") from exc
    try:
        return _sheet_rows(workbook, sheet_name)[1:]
    finally:
        workbook.close()


def _read_row(
    raw_row: Sequence[object],
    cached_row: Sequence[object] | None,
    index_of: dict[str, int],
) -> tuple[dict[str, object], dict[str, str]]:
    """按 13 列取出可用值；第二个返回值记录不可读的列与原因。"""
    values: dict[str, object] = {}
    unreadable: dict[str, str] = {}

    for column in INPUT_COLUMNS:
        position = index_of[column]
        raw_value = _at(raw_row, position)

        if is_formula_cell(raw_value):
            cached = _at(cached_row, position) if cached_row is not None else None
            if cached is None:
                unreadable[column] = "公式单元格没有缓存值（文件未经 Excel 计算保存）"
                values[column] = None
            elif is_error_cell(cached):
                unreadable[column] = f"公式结果为 Excel 错误值 {str(cached).strip()}"
                values[column] = None
            else:
                values[column] = cached
        elif is_error_cell(raw_value):
            unreadable[column] = f"单元格为 Excel 错误值 {str(raw_value).strip()}"
            values[column] = None
        else:
            values[column] = raw_value

    return values, unreadable


def _input_error_reasons(record: QARecord, unreadable: dict[str, str]) -> list[str]:
    """一条记录的输入失败原因；空列表表示该记录有效。

    - 编号、q、a 缺失或不可读 → 输入失败；
    - ref 为空合法，但**不可读**的 ref 不能当作资料文本，同样记输入失败。
    """
    reasons: list[str] = []
    for column in REQUIRED_TEXT_COLUMNS:
        if column in unreadable:
            reasons.append(f"{column}：不可读取（{unreadable[column]}）")
        elif _required_text(record, column) is None:
            reasons.append(f"{column}：缺失或为空")

    for field in REF_FIELDS:
        if field in unreadable:
            reasons.append(f"{field}：不可读取（{unreadable[field]}）")
    return reasons


def _required_text(record: QARecord, column: str) -> str | None:
    """取必填列的已归一化文本。"""
    if column == ID_COLUMN:
        return record.record_id
    if column == Q_COLUMN:
        return record.q
    return record.a


def _batch_blockers(
    records: Sequence[QARecord], counts: PrecheckCounts, limits: LimitsConfig
) -> tuple[Blocker, ...]:
    """批次级问题：编号重复、超记录数上限、没有有效记录。"""
    blockers: list[Blocker] = []

    seen: dict[str, list[int]] = {}
    for record in records:
        if record.record_id is not None:
            seen.setdefault(record.record_id, []).append(record.source_row)
    duplicates = {key: rows for key, rows in seen.items() if len(rows) > 1}
    if duplicates:
        detail = "；".join(
            f"{key}（来源行 {', '.join(str(row) for row in rows)}）"
            for key, rows in sorted(duplicates.items())
        )
        blockers.append(
            Blocker(
                code=BLOCKER_DUPLICATE_RECORD_ID,
                message=f"编号必须唯一，发现重复编号：{detail}",
                field=",".join(sorted(duplicates)),
                source_rows=tuple(
                    row for rows in duplicates.values() for row in rows
                ),
            )
        )

    if counts.total is not None and counts.total > limits.max_records:
        blockers.append(
            Blocker(
                code=BLOCKER_TOO_MANY_RECORDS,
                message=(
                    f"非空记录 {counts.total} 条超过上限 {limits.max_records} 条；"
                    "不截取输入，请分批处理"
                ),
                field="record_count",
            )
        )

    if counts.valid == 0:
        blockers.append(
            Blocker(
                code=BLOCKER_NO_VALID_RECORDS,
                message="没有可处理的记录：全部行都属空行或输入失败",
                field="records",
            )
        )

    return tuple(blockers)


def _warnings(header_text: Sequence[str], numeric_id_rows: Sequence[int]) -> tuple[PrecheckWarning, ...]:
    warnings: list[PrecheckWarning] = []

    extra = sorted({name for name in header_text if name and name not in INPUT_COLUMNS})
    if extra:
        warnings.append(
            PrecheckWarning(
                code=WARNING_EXTRA_COLUMNS,
                message=f"忽略非必需列，不发送给模型：{'、'.join(extra)}",
            )
        )

    if numeric_id_rows:
        rows = "、".join(str(row) for row in numeric_id_rows[:5])
        more = "" if len(numeric_id_rows) <= 5 else f" 等 {len(numeric_id_rows)} 行"
        warnings.append(
            PrecheckWarning(
                code=WARNING_NUMERIC_RECORD_ID,
                message=(
                    f"编号列为数值型（来源行 {rows}{more}），Excel 可能已丢弃前导零；"
                    "编号仍按原值关联，未补零"
                ),
            )
        )

    return tuple(warnings)


def _blocked(
    *,
    sheet_name: str | None,
    digest: str,
    size_bytes: int,
    counts: PrecheckCounts,
    blockers: tuple[Blocker, ...],
) -> ParsedInput:
    return ParsedInput(
        report=PrecheckReport(
            sheet_name=sheet_name,
            status="blocked",
            file_sha256=digest,
            file_size_bytes=size_bytes,
            counts=counts,
            blockers=blockers,
        ),
        records=(),
    )
