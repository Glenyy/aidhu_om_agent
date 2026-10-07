"""S02-05：合成 Excel 样例生成器。

仓库不提交二进制 .xlsx；测试与手动审阅都在运行时按需生成。
每个生成器返回写出的路径，并只构造**当前场景**需要的异常，便于逐项核对。
"""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from pathlib import Path

from openpyxl import Workbook

from aidhu_om_agent.excel.reader import PREFERRED_SHEET
from aidhu_om_agent.schemas.qa import A_COLUMN, ID_COLUMN, INPUT_COLUMNS, Q_COLUMN, REF_FIELDS

HEADER: tuple[str, ...] = INPUT_COLUMNS
OTHER_SHEET = "Sheet1"

Q_TEXT = "如何办理校园卡？"
A_TEXT = "请到校园卡中心办理。"


def row(cells: dict[str, object]) -> list[object]:
    """按列名构造一行；未给出的列为 None。"""
    return [cells.get(column) for column in INPUT_COLUMNS]


def qa_row(
    record_id: object,
    q: object = Q_TEXT,
    a: object = A_TEXT,
    refs: Sequence[object] = (),
) -> list[object]:
    """构造常规数据行；``refs`` 依次填入 ref1、ref2……"""
    cells: dict[str, object] = {ID_COLUMN: record_id, Q_COLUMN: q, A_COLUMN: a}
    for field, text in zip(REF_FIELDS, refs):
        cells[field] = text
    return row(cells)


def write_workbook(
    path: Path,
    rows: Iterable[Sequence[object]],
    *,
    header: Sequence[object] = HEADER,
    sheet_name: str = PREFERRED_SHEET,
    extra_sheets: Iterable[tuple[str, Sequence[Sequence[object]]]] = (),
) -> Path:
    """写出一个合成工作簿并返回路径。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(list(header))
    for data_row in rows:
        sheet.append(list(data_row))

    for name, extra_rows in extra_sheets:
        extra = workbook.create_sheet(title=name)
        extra.append(list(header))
        for data_row in extra_rows:
            extra.append(list(data_row))

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


# ---------------------------------------------------------------- 正常场景


def normal_workbook(path: Path) -> Path:
    """3 条有效记录，ref 分别有值、全空、多行文本；末尾 1 条空行。

    多行文本用 ``\\n``：openpyxl 在 Windows 写入时会把 ``\\n`` 再转成 ``\\r\\n``，
    因此在样例里直接写 ``\\r\\n`` 会得到 ``\\r\\r\\n``（解析后成双换行）。
    真实 Excel 文件存的是 LF，CRLF 归一化另由 `normalize_text` 的纯函数单测覆盖。
    """
    return write_workbook(
        path,
        [
            qa_row("1", "问题一", "回答一", ["资料一"]),
            qa_row("2", "问题二", "回答二", ()),
            qa_row("A-003", "问题三", "回答三", ["资料三第一行\n资料三第二行"]),
            row({}),
        ],
    )


def numeric_id_workbook(path: Path) -> Path:
    """编号列为数值型：7、8。"""
    return write_workbook(path, [qa_row(7), qa_row(8, "问题二", "回答二", ["资料二"])])


def text_id_leading_zero_workbook(path: Path) -> Path:
    """编号列为文本型且带前导零：``"007"``。"""
    return write_workbook(path, [qa_row("007"), qa_row("008", "问题二")])


# ---------------------------------------------------------------- 批次级异常


def missing_column_workbook(path: Path, column: str = REF_FIELDS[-1]) -> Path:
    """缺少一列必要列。"""
    header = [name for name in HEADER if name != column]
    return write_workbook(path, [qa_row("1")], header=header)


def duplicated_header_workbook(path: Path, column: str = Q_COLUMN) -> Path:
    """表头中同一列名出现两次。"""
    header = list(HEADER) + [column]
    return write_workbook(path, [qa_row("1")], header=header)


def duplicated_id_workbook(path: Path) -> Path:
    """两条来源行使用同一编号。"""
    return write_workbook(
        path,
        [qa_row("1", "问题一"), qa_row("2", "问题二"), qa_row("1", "问题三")],
    )


def multi_sheet_workbook(path: Path) -> Path:
    """没有 QA_REF，且有两个可见工作表。"""
    return write_workbook(
        path,
        [qa_row("1")],
        sheet_name=OTHER_SHEET,
        extra_sheets=[("Another", [qa_row("2")])],
    )


def single_sheet_workbook(path: Path) -> Path:
    """没有 QA_REF，但只有一张可见表：应自动选中。"""
    return write_workbook(path, [qa_row("1")], sheet_name=OTHER_SHEET)


def extra_column_workbook(path: Path) -> Path:
    """含非必需列。"""
    header = list(HEADER) + ["人工标签"]
    return write_workbook(path, [qa_row("1") + ["回答正确"]], header=header)


def oversized_workbook(path: Path, count: int) -> Path:
    """``count`` 条非空记录，用于配合较小的 max_records 触发上限。"""
    return write_workbook(path, [qa_row(str(index)) for index in range(1, count + 1)])


# ------------------------------------------------- S06-02 的规模防护（三种各一）


def heavy_text_workbook(path: Path, *, rows: int = 100) -> Path:
    """压缩后十几 KB、**解压后约 3 MB** 的工作簿（解压总量防护用）。

    每行 q 都是接近 Excel 单元格上限的长文本，且逐行不同（各不相同才会被
    openpyxl 各自写进 sharedStrings，不会被去重成一份）。
    """
    return write_workbook(
        path,
        [qa_row(str(index), "A" * 32760 + str(index)) for index in range(rows)],
    )


def many_cells_workbook(path: Path, *, row: int = 1000, column: int = 600) -> Path:
    """声明维度达 ``row × column`` 的工作簿（单元格数防护用）。

    只写了一个远端单元格，文件本身仍只有几 KB——真实数据都在表头附近，
    但**工作表声明的尺寸**是 60 万个单元格。这正是要挡住的那类文件。
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = PREFERRED_SHEET
    sheet.append(list(HEADER))
    sheet.append(qa_row("1", Q_TEXT, A_TEXT, ["资料一"]))
    sheet.cell(row=row, column=column).value = "远端标记"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def many_sheets_workbook(path: Path, *, sheets: int = 51) -> Path:
    """``sheets`` 张**可见**工作表（可见工作表数防护用）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = PREFERRED_SHEET
    sheet.append(list(HEADER))
    sheet.append(qa_row("1", Q_TEXT, A_TEXT, ["资料一"]))
    for index in range(1, sheets):
        extra = workbook.create_sheet(title=f"EXTRA_{index}")
        extra.append(list(HEADER))

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


# ---------------------------------------------------------------- 单条异常


def partial_failure_workbook(path: Path) -> Path:
    """缺编号、缺 q、缺 a 各一条，外加有效记录与全空行。"""
    return write_workbook(
        path,
        [
            qa_row("1", "问题一", "回答一", ["资料一"]),
            qa_row(None, "问题二", "回答二"),
            qa_row("3", None, "回答三"),
            qa_row("4", "问题四", None),
            row({}),
            row({column: "   " for column in INPUT_COLUMNS}),
        ],
    )


def empty_refs_workbook(path: Path) -> Path:
    """ref1—ref10 全为空：仍是有效记录（资料不足属业务判断）。"""
    return write_workbook(path, [qa_row("1", "问题一", "回答一", ())])


def error_value_workbook(path: Path) -> Path:
    """第 2 条的 q 列为 Excel 错误值、ref1 列也是错误值；第 1 条正常。"""
    cells: dict[str, object] = {
        ID_COLUMN: "2",
        Q_COLUMN: "#N/A",
        A_COLUMN: A_TEXT,
        "ref1": "#VALUE!",
    }
    return write_workbook(path, [qa_row("1", "问题一", "回答一", ["资料一"]), row(cells)])


def formula_without_cache_workbook(path: Path) -> Path:
    """第 2 条的 q 为公式且未保存缓存值（openpyxl 写出的文件没有缓存结果）。"""
    cells: dict[str, object] = {
        ID_COLUMN: "2",
        Q_COLUMN: '=CONCATENATE("问题","二")',
        A_COLUMN: A_TEXT,
        "ref1": "资料二",
    }
    return write_workbook(path, [qa_row("1", "问题一", "回答一", ["资料一"]), row(cells)])


def all_invalid_workbook(path: Path) -> Path:
    """没有任何有效记录：只有缺 q 的行和空行。"""
    return write_workbook(path, [qa_row("1", None, "回答一"), row({})])
