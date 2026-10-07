"""S05-02 导出快照与两份产物装配的单测。

覆盖：分类结果表一条一行且按 order_index 排序、标签只在已分类记录上出现、
复核清单只放被标记记录且两个人工列留空、失败清单列出阶段与可重试性、
概况表按固定标签顺序给出计数与「未处理」、JSONL 每条一行且原文不截断、
阶段二更正不覆盖阶段一、长文本截断有账可查、两份文件成对写出、半组不发布。

数据由**模拟客户端**跑出来，全程零真实调用。
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from aidhu_om_agent.agent.mock_samples import SCENARIO_FORCED_FAILURE, MockClient
from aidhu_om_agent.repositories import exports as exports_repo
from aidhu_om_agent.repositories import jobs as jobs_repo
from aidhu_om_agent.repositories import records as records_repo
from aidhu_om_agent.repositories import runs as runs_repo
from aidhu_om_agent.schemas.export import (
    ARTIFACT_EXCEL,
    ARTIFACT_JSONL,
    CLASSIFICATION_COLUMNS,
    EXCEL_CELL_MAX_CHARS,
    EXCEL_DIGEST_NOTE,
    EXPORT_CONTRACT_VERSION,
    FAILURE_COLUMNS,
    JSONL_RECORD_FIELDS,
    REVIEW_COLUMNS,
    SHEET_CLASSIFICATION,
    SHEET_FAILURES,
    SHEET_REVIEW,
    SHEET_SUMMARY,
    SUMMARY_FIELDS,
    TEXT_GUARD_PREFIX,
    TRUNCATION_MARKER,
    download_name,
)
from aidhu_om_agent.schemas.qa import REF_FIELDS
from aidhu_om_agent.services.batches import counts_of, create_run
from aidhu_om_agent.services.exports import (
    ExportStateError,
    TEMP_SUFFIX,
    WrittenArtifact,
    build_document,
    load_snapshot,
    publish_export,
    write_document,
)
from aidhu_om_agent.storage import read_transaction, write_transaction
from aidhu_om_agent.worker import Worker
from fixtures.excel_samples import qa_row, write_workbook
from test_batches import make_config, prepare_validation

STAMP = "2026-10-07T00:00:00+00:00"
EXPORT_ID = "e" * 32


# ------------------------------------------------------------------ 夹具


def mixed_workbook(path: Path) -> Path:
    """5 条记录：两条正常、一条强制复核、一条技术失败、一条缺 a（输入失败）。"""
    return write_workbook(
        path,
        [
            qa_row("1", "问题一", "回答一", ["资料一"]),
            qa_row("2", "问题二", "回答二", ["资料二"]),
            qa_row("3", "问题三", "回答三", ["资料三"]),
            qa_row("4", "问题四", None, ["资料四"]),
            qa_row("5", "问题五", "回答五", ["资料五"]),
        ],
    )


def mixed_factory(record: records_repo.RecordRow) -> MockClient:
    """按编号指定情境：2 强制复核、3 技术失败、5 重大更正，其余回答正确。"""
    scenario = "correct"
    if record.record_id == "2":
        scenario = "forced_review"
    elif record.record_id == "3":
        scenario = SCENARIO_FORCED_FAILURE
    elif record.record_id == "5":
        scenario = "correction"
    return MockClient(record.to_qa_record(), scenario=scenario)


def prepare_run(config, workbook: Path, tmp_path: Path):
    """走「上传 → 预检 → 建批次」，返回 (database, created)。"""
    database, validation_id = prepare_validation(tmp_path, workbook, config)
    return database, create_run(database, config, validation_id=validation_id)


def run_worker_once(config, factory) -> None:
    worker = Worker(
        config, poll_seconds=0.01, sleep=lambda _: None, client_factory=factory
    )
    worker.start()
    try:
        worker.run_once()
    finally:
        worker.stop()


def snapshot_of(database, run_id: str, *, captured_at: str = STAMP):
    """在**一个读事务**里装载快照；计数与界面共用 `counts_of`。"""
    connection = database.connect()
    try:
        with read_transaction(connection) as snapshot_connection:
            run = runs_repo.get_run(snapshot_connection, run_id)
            counts = counts_of(
                run,
                records_repo.count_by_status(snapshot_connection, run_id),
                records_repo.count_review_required(snapshot_connection, run_id),
            )
            return load_snapshot(
                snapshot_connection, run_id, captured_at=captured_at, counts=counts
            )
    finally:
        connection.close()


def sheet_rows(document, name: str) -> list[tuple]:
    sheet = next(item for item in document.sheets if item.name == name)
    return list(sheet.rows)


def jsonl_records(document) -> list[dict]:
    text = document.jsonl_bytes.decode("utf-8")
    return [json.loads(line) for line in text.splitlines() if line]


@pytest.fixture
def mixed_run(tmp_path: Path):
    """跑完一批混合记录，返回 (config, database, run_id)。"""
    config = make_config(tmp_path)
    workbook = mixed_workbook(tmp_path / "src" / "mixed.xlsx")
    database, created = prepare_run(config, workbook, tmp_path)
    run_worker_once(config, mixed_factory)
    return config, database, created.run_id


# ------------------------------------------------------- 分类结果表（S05-02）


def test_classification_sheet_has_one_row_per_record_in_input_order(mixed_run) -> None:
    _, database, run_id = mixed_run
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)
    rows = sheet_rows(document, SHEET_CLASSIFICATION)

    assert len(rows) == 5
    # 序号 1..N 由 order_index 生成；来源行含标题行，所以是 2..6。
    assert [row[0] for row in rows] == [1, 2, 3, 4, 5]
    assert [row[2] for row in rows] == [2, 3, 4, 5, 6]
    assert [row[1] for row in rows] == ["1", "2", "3", "4", "5"]


def test_labels_appear_only_on_classified_records(mixed_run) -> None:
    _, database, run_id = mixed_run
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)
    by_number = {row[0]: row for row in sheet_rows(document, SHEET_CLASSIFICATION)}

    assert by_number[1][3] == "回答正确"
    assert by_number[1][4]  # 理由非空
    assert by_number[1][7] == "已分类"
    assert by_number[2][6] == "是"  # 强制复核
    assert by_number[3][7] == "技术失败"
    assert by_number[4][7] == "输入失败"
    # 失败与输入失败的行**留空**，不填造标签与理由。
    assert by_number[3][3] == "" and by_number[3][4] == ""
    assert by_number[4][3] == "" and by_number[4][4] == ""


def test_evidence_column_names_the_ref_and_quotes_it(mixed_run) -> None:
    _, database, run_id = mixed_run
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)
    first = sheet_rows(document, SHEET_CLASSIFICATION)[0]
    ref_id, quote = first[5].split("：", 1)
    assert ref_id in REF_FIELDS
    assert quote and quote in "资料一"


# --------------------------------------------------------- 复核清单（S05-02）


def test_review_sheet_carries_only_flagged_records_with_blank_human_columns(
    mixed_run,
) -> None:
    _, database, run_id = mixed_run
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)
    rows = sheet_rows(document, SHEET_REVIEW)

    assert len(document.sheets[1].columns) == len(REVIEW_COLUMNS)
    # 只有 2 号（强制复核）与 5 号（阶段二重大更正）被标记。
    assert [row[0] for row in rows] == [2, 5]
    row = rows[0]
    assert row[0] == 2  # 序号仍是它在整批里的位置，不重新编号
    assert row[1] == "2"
    assert row[3] == "问题二"  # q
    assert row[4] == "回答二"  # a
    # 十个 ref 列齐备：有值的第一列是资料，其余留空。
    refs = row[5 : 5 + len(REF_FIELDS)]
    assert refs[0] == "资料二"
    assert set(refs[1:]) == {""}
    # 两个人工填写列导出时留空，agent 不预填、也不回读覆盖。
    assert row[-2] == "" and row[-1] == ""


# --------------------------------------------------------- 失败清单（S05-02）


def test_failure_sheet_reports_stage_code_and_retryability(mixed_run) -> None:
    _, database, run_id = mixed_run
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)
    rows = sheet_rows(document, SHEET_FAILURES)

    assert len(rows) == 1
    row = rows[0]
    assert row[0] == 3
    assert row[1] == "3"
    assert row[2] == 4
    assert row[3] == "阶段二"
    assert row[4]  # 错误码非空
    assert row[5]  # 错误信息非空
    assert row[6] == "3"  # 三次尝试全部失败
    assert row[7] == "否"


# --------------------------------------------------------- 运行概况（S05-02）


def test_summary_lists_every_required_field_in_order(mixed_run) -> None:
    _, database, run_id = mixed_run
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)
    rows = sheet_rows(document, SHEET_SUMMARY)

    assert [row[0] for row in rows] == list(SUMMARY_FIELDS)
    values = {row[0]: row[1] for row in rows}
    assert values["批次标识"] == run_id
    assert values["总数"] == 5
    assert values["有效记录"] == 4
    assert values["输入失败"] == 1
    assert values["已分类"] == 3
    assert values["技术失败"] == 1
    assert values["未处理"] == 0
    assert values["需复核"] == 2
    assert values["批次状态（捕获时）"] == "部分失败（partial_failed）"
    assert values["被截断单元格数"] == 0
    assert values["文本前缀防护单元格数"] == 0
    assert values["捕获时间"] == STAMP
    assert values["捕获修订"] == snapshot_of(database, run_id).run.revision
    # 两份产物的事实都已写进概况表，可回 artifacts 表逐字核对。
    assert values["导出合同版本"] == EXPORT_CONTRACT_VERSION
    assert values["JSONL 文件"] == download_name(
        ARTIFACT_JSONL, run_id=run_id, export_id=EXPORT_ID
    )
    assert values["JSONL 文件摘要"] == document.jsonl_sha256
    assert values["Excel 文件摘要"] == EXCEL_DIGEST_NOTE


def test_summary_counts_unprocessed_records_for_an_unfinished_batch(tmp_path) -> None:
    """没跑过 worker 的批次照样能导出，但必须自己说明「还没处理完」。"""
    config = make_config(tmp_path)
    workbook = mixed_workbook(tmp_path / "src" / "pending.xlsx")
    database, created = prepare_run(config, workbook, tmp_path)

    snapshot = snapshot_of(database, created.run_id)
    document = build_document(snapshot, export_id=EXPORT_ID)
    values = {row[0]: row[1] for row in sheet_rows(document, SHEET_SUMMARY)}

    assert snapshot.unprocessed == 4
    assert values["未处理"] == 4
    assert values["已分类"] == 0
    assert values["批次状态（捕获时）"] == "排队中（queued）"
    assert [row[7] for row in sheet_rows(document, SHEET_CLASSIFICATION)] == [
        "未处理",
        "未处理",
        "未处理",
        "输入失败",
        "未处理",
    ]


# ------------------------------------------------------------- JSONL（S05-02）


def test_jsonl_has_one_line_per_record_with_ordinals_and_full_refs(mixed_run) -> None:
    _, database, run_id = mixed_run
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)
    records = jsonl_records(document)

    assert len(records) == 5
    assert [item["ordinal"] for item in records] == [1, 2, 3, 4, 5]
    assert [item["order_index"] for item in records] == [0, 1, 2, 3, 4]
    assert all(item["export_id"] == EXPORT_ID for item in records)
    assert all(item["run_id"] == run_id for item in records)
    assert all(list(item["refs"]) == list(REF_FIELDS) for item in records)
    assert records[1]["refs"]["ref1"] == "资料二"
    # 十个 ref 列都在；没填的留空（null），不补空字符串冒充有值。
    assert records[1]["refs"]["ref2"] is None
    assert records[0]["versions"]["export_contract"] == EXPORT_CONTRACT_VERSION
    # 技术失败的记录：阶段一一次就过，阶段二三次全被拒；三次失败都记在账上。
    failed_attempts = records[2]["attempt_summary"]
    assert failed_attempts == {
        "total": 4,
        "stage1": 1,
        "stage2": 3,
        "failed": 3,
        "unknown_after_interrupt": 0,
        "simulated": 4,
    }
    assert records[2]["failure"]["attempt_count"] == 3
    # 正常记录：两阶段各一次。
    assert records[0]["attempt_summary"]["total"] == 2
    assert records[0]["attempt_summary"]["failed"] == 0


def test_jsonl_never_carries_rejected_model_output(mixed_run) -> None:
    """被拒的模型输出**只在记录详情**暴露，导出不承接（S06 阶段文档 §3）。

    这批数据里确实有 3 次被校验拒绝的阶段二尝试（`call_attempts.final_content`
    非空），所以这里排的是「导出层漏带」，不是「本来就没有」。
    """
    _, database, run_id = mixed_run
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)
    records = jsonl_records(document)

    failed = next(item for item in records if item["failure"])
    assert failed["failure"]["attempt_count"] == 3
    for item in records:
        assert list(item) == list(JSONL_RECORD_FIELDS)
        assert not any("raw_output" in key or "content" in key for key in item)
    # 拒过的正文还在库里（导出没有它，不等于没有发生）。
    with read_transaction(database.connect()) as snapshot:
        stored = snapshot.execute(
            "SELECT COUNT(*) FROM call_attempts WHERE final_content IS NOT NULL"
        ).fetchone()[0]
    assert stored >= failed["failure"]["attempt_count"] > 0


def test_jsonl_keeps_null_labels_and_blank_input_for_unfinished_records(
    mixed_run,
) -> None:
    _, database, run_id = mixed_run
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)
    records = jsonl_records(document)

    assert records[0]["label"] == "回答正确"
    assert records[0]["review_required"] is False
    assert records[1]["review_required"] is True
    assert records[2]["label"] is None and records[2]["reason"] is None
    assert records[2]["status"] == "failed"
    assert records[2]["failure"]["code"] and records[2]["failure"]["retryable"] is False
    # 输入失败的记录：原文照留，只是没被判过；标签为空、阶段结果为空。
    assert records[3]["status"] == "input_invalid"
    assert records[3]["q"] == "问题四" and records[3]["a"] is None
    assert records[3]["label"] is None
    assert records[3]["stage1"] is None and records[3]["stage2"] is None


def test_jsonl_keeps_stage1_alongside_a_stage2_correction(mixed_run) -> None:
    """阶段二的更正**不覆盖**阶段一原结果：两者在同一行里各就各位。"""
    _, database, run_id = mixed_run
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)
    corrected = jsonl_records(document)[4]

    assert corrected["stage1"]["required_points"]
    assert corrected["stage2"]["stage1_corrections"]
    assert corrected["stage2"]["review_required"] is True


# ------------------------------------------------------- 长文本与防护（S05-02）


def forced_factory(scenario: str):
    """把一整批都按同一情境跑；用于把记录推到需要复核的表里。"""

    def factory(record: records_repo.RecordRow) -> MockClient:
        return MockClient(record.to_qa_record(), scenario=scenario)

    return factory


def single_record_run(tmp_path: Path, refs):
    """一条记录、指定 ref 的小批次；走强制复核，使 ref 落进复核清单。"""
    config = make_config(tmp_path)
    workbook = write_workbook(
        tmp_path / "src" / "one.xlsx", [qa_row("1", "问题", "回答", refs)]
    )
    database, created = prepare_run(config, workbook, tmp_path)
    run_worker_once(config, forced_factory("forced_review"))
    return config, database, created.run_id


def inject_refs(database, run_id: str, refs: dict[str, str]) -> None:
    """直接改库里的 `refs_json`。

    有些值**造不进 .xlsx**：openpyxl 会把 `=` 开头的字符串当成公式写出去（读回时
    被预检按「无缓存值的公式」拒收），而超过 32767 字符的单元格 openpyxl 会**静默
    截断**（Excel 自己的上限）。导出层读的是库，所以从这里注入才能验证它不假定
    「上游一定干净」——顺带说明这类防护是纵深防御，不是唯一的拦截点。
    """
    connection = database.connect()
    try:
        with write_transaction(connection) as tx:
            row = tx.execute(
                "SELECT record_key FROM records WHERE run_id = ?", (run_id,)
            ).fetchone()
            tx.execute(
                "UPDATE records SET refs_json = ? WHERE record_key = ?",
                (records_repo.refs_to_json(refs), row["record_key"]),
            )
    finally:
        connection.close()


def test_long_ref_is_truncated_in_excel_but_complete_in_jsonl(tmp_path: Path) -> None:
    """超出单元格上限时截断并记账；原文在 JSONL 里一字不少。"""
    long_text = "长" * (EXCEL_CELL_MAX_CHARS + 500)
    _config, database, run_id = single_record_run(tmp_path, ["干净资料"])
    inject_refs(database, run_id, {"ref1": long_text})
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)

    cell = sheet_rows(document, SHEET_REVIEW)[0][5]  # 复核清单的 ref1
    assert len(cell) == EXCEL_CELL_MAX_CHARS
    assert cell.startswith(long_text[:100])
    assert cell.endswith(TRUNCATION_MARKER)

    values = {row[0]: row[1] for row in sheet_rows(document, SHEET_SUMMARY)}
    assert values["被截断单元格数"] == 1
    assert values["被截断单元格位置"] == (
        f"复核清单!ref1 第 2 行（原 {len(long_text)} 字符）"
    )
    # 没有清理过控制字符时留空，不写「无」之类的占位文字。
    assert values["清理控制字符位置"] == ""

    record = jsonl_records(document)[0]
    assert record["refs"]["ref1"] == long_text  # JSONL 不截断，一字不少


def test_dash_and_at_refs_from_the_workbook_are_prefixed(tmp_path: Path) -> None:
    """`-`/`@` 开头的资料**能从 .xlsx 正常读进来**，导出时必须加文本前缀。"""
    _config, database, run_id = single_record_run(tmp_path, ["-2+3", "@SUM(A1)"])
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)

    row = sheet_rows(document, SHEET_REVIEW)[0]
    assert row[5] == TEXT_GUARD_PREFIX + "-2+3"
    assert row[6] == TEXT_GUARD_PREFIX + "@SUM(A1)"
    values = {item[0]: item[1] for item in sheet_rows(document, SHEET_SUMMARY)}
    assert values["文本前缀防护单元格数"] == 2
    # 前缀只是 Excel 落盘的防护，不改数据：JSONL 里是没有前缀的原文。
    record = jsonl_records(document)[0]
    assert record["refs"]["ref1"] == "-2+3"
    assert record["refs"]["ref2"] == "@SUM(A1)"


def test_control_characters_are_cleaned_in_excel_and_counted(tmp_path: Path) -> None:
    """Excel 拒绝的控制字符被删掉并记账；原文仍在 JSONL 里。"""
    dirty = "第一行\x0b第二行\x00尾"
    _config, database, run_id = single_record_run(tmp_path, ["干净资料"])
    inject_refs(database, run_id, {"ref1": dirty, "ref2": "=1+1"})
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)

    row = sheet_rows(document, SHEET_REVIEW)[0]
    assert row[5] == "第一行第二行尾"
    # 同一个单元格里两种改写可以叠加：先删控制字符，再按清理后的首字符判断防护。
    assert row[6] == TEXT_GUARD_PREFIX + "=1+1"
    values = {row[0]: row[1] for row in sheet_rows(document, SHEET_SUMMARY)}
    assert values["清理控制字符单元格数"] == 1
    assert values["清理控制字符位置"] == "复核清单!ref1 第 2 行（清理 2 个控制字符）"
    assert values["文本前缀防护单元格数"] == 1
    record = jsonl_records(document)[0]
    assert record["refs"]["ref1"] == dirty
    assert record["refs"]["ref2"] == "=1+1"


# ------------------------------------------------------------- 写出（S05-02）


def test_write_document_creates_both_files_with_digests(
    mixed_run, tmp_path: Path
) -> None:
    config, database, run_id = mixed_run
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)
    artifacts = write_document(document, outputs_root=config.paths.outputs)

    assert [artifact.kind for artifact in artifacts] == [
        ARTIFACT_EXCEL,
        ARTIFACT_JSONL,
    ]
    for artifact in artifacts:
        assert artifact.path.is_file()
        assert artifact.size_bytes == artifact.path.stat().st_size > 0
        assert artifact.relative_path.startswith(f"{run_id}/{EXPORT_ID}/")
        assert artifact.path.parent.name == EXPORT_ID
    jsonl = next(a for a in artifacts if a.kind == ARTIFACT_JSONL)
    assert jsonl.path.read_bytes() == document.jsonl_bytes

    workbook = load_workbook(next(a for a in artifacts if a.kind == ARTIFACT_EXCEL).path)
    assert workbook.sheetnames == [
        "分类结果",
        "复核清单",
        "失败清单",
        "运行概况",
    ]
    assert workbook[SHEET_CLASSIFICATION].max_row == 6  # 表头 + 5 条
    sheet = workbook[SHEET_SUMMARY]
    assert [sheet.cell(row=index, column=1).value for index in (1, 2)] == [
        "项目",
        "批次标识",
    ]


def test_write_document_leaves_no_temp_files_behind(mixed_run) -> None:
    """临时文件只用于「写好再改名」；成功与失败都不在导出目录里留下 `.part`。"""
    config, database, run_id = mixed_run
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)
    artifacts = write_document(document, outputs_root=config.paths.outputs)

    directory = config.paths.outputs / run_id / EXPORT_ID
    assert list(directory.glob(f"*{TEMP_SUFFIX}")) == []
    assert sorted(item.name for item in directory.iterdir()) == sorted(
        artifact.path.name for artifact in artifacts
    )


def test_written_columns_match_the_declared_headers(mixed_run) -> None:
    config, database, run_id = mixed_run
    document = build_document(snapshot_of(database, run_id), export_id=EXPORT_ID)
    artifacts = write_document(document, outputs_root=config.paths.outputs)
    workbook = load_workbook(next(a for a in artifacts if a.kind == ARTIFACT_EXCEL).path)

    assert [
        workbook[SHEET_CLASSIFICATION].cell(row=1, column=index + 1).value
        for index in range(len(CLASSIFICATION_COLUMNS))
    ] == list(CLASSIFICATION_COLUMNS)
    assert [
        workbook[SHEET_REVIEW].cell(row=1, column=index + 1).value
        for index in range(len(REVIEW_COLUMNS))
    ] == list(REVIEW_COLUMNS)
    assert [
        workbook[SHEET_FAILURES].cell(row=1, column=index + 1).value
        for index in range(len(FAILURE_COLUMNS))
    ] == list(FAILURE_COLUMNS)


# ------------------------------------------------------------- 发布（S05-03）


def test_publish_export_refuses_a_half_group(mixed_run) -> None:
    _, database, run_id = mixed_run
    job_id = jobs_repo.new_job_id()
    export_id = exports_repo.new_export_id()
    connection = database.connect()
    try:
        with write_transaction(connection) as tx:
            jobs_repo.insert_job(
                tx,
                job_id=job_id,
                run_id=run_id,
                kind="export",
                mode="manual",
                payload={},
                created_at=STAMP,
            )
            exports_repo.insert_export(
                tx,
                export_id=export_id,
                job_id=job_id,
                run_id=run_id,
                source="manual",
                scheduled_revision=0,
                created_at=STAMP,
            )
    finally:
        connection.close()

    only_excel = WrittenArtifact(
        kind=ARTIFACT_EXCEL,
        path=Path("x.xlsx"),
        relative_path="x.xlsx",
        download_name="x.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        size_bytes=1,
        sha256="0" * 64,
    )
    connection = database.connect()
    try:
        with pytest.raises(ExportStateError, match="不成对"):
            with write_transaction(connection) as tx:
                publish_export(
                    tx,
                    job_id=job_id,
                    export_id=export_id,
                    artifacts=[only_excel],
                    created_at=STAMP,
                )
    finally:
        connection.close()

    # 事务回滚：既没有登记半份文件，任务也没有被标成 completed。
    connection = database.connect()
    try:
        assert exports_repo.list_artifacts(connection, export_id) == ()
        assert jobs_repo.get_job(connection, job_id).status == "queued"
    finally:
        connection.close()


def test_publish_export_registers_both_files_and_completes_the_job(mixed_run) -> None:
    config, database, run_id = mixed_run
    job_id = jobs_repo.new_job_id()
    export_id = exports_repo.new_export_id()
    connection = database.connect()
    try:
        with write_transaction(connection) as tx:
            jobs_repo.insert_job(
                tx,
                job_id=job_id,
                run_id=run_id,
                kind="export",
                mode="manual",
                payload={},
                created_at=STAMP,
            )
            exports_repo.insert_export(
                tx,
                export_id=export_id,
                job_id=job_id,
                run_id=run_id,
                source="manual",
                scheduled_revision=0,
                created_at=STAMP,
            )
    finally:
        connection.close()

    document = build_document(snapshot_of(database, run_id), export_id=export_id)
    artifacts = write_document(document, outputs_root=config.paths.outputs)
    connection = database.connect()
    try:
        with write_transaction(connection) as tx:
            rows = publish_export(
                tx,
                job_id=job_id,
                export_id=export_id,
                artifacts=artifacts,
                created_at=STAMP,
            )
        job = jobs_repo.get_job(connection, job_id)
        assert [row.kind for row in rows] == [ARTIFACT_EXCEL, ARTIFACT_JSONL]
        assert job.status == "completed"
        assert job.result["export_id"] == export_id
        assert len(job.result["artifacts"]) == 2
    finally:
        connection.close()
