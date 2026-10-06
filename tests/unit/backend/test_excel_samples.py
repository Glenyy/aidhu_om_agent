"""S03-07：合成样例定义的单测。

样例是界面手动验证的输入来源（仓库不提交二进制 .xlsx）。这里守住的是**界面里
能点到什么**：每份样例的预检结果必须与它对外宣称的说明一致，且编号 1/2/3/6/9
在模拟模式下必须覆盖全部五类情境——否则手动指南里写的结果就对不上。

全部为合成数据，不调用模型、不访问网络。
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from aidhu_om_agent.agent.mock_samples import SCENARIOS, scenario_for
from aidhu_om_agent.excel.reader import PREFERRED_SHEET, precheck
from aidhu_om_agent.excel.samples import SAMPLES, build_workbook_bytes, get_sample
from aidhu_om_agent.schemas.qa import INPUT_COLUMNS

#: 每份样例宣称的预检结果；与 description 一一对应。
EXPECTED_OUTCOME: dict[str, dict[str, int | str]] = {
    "five-scenarios": {
        "status": "passed",
        "total": 5,
        "valid": 5,
        "input_invalid": 0,
        "skipped_blank_rows": 1,
        "row_errors": 0,
        "valid_records": 5,
    },
    "no-refs": {
        "status": "passed",
        "total": 1,
        "valid": 1,
        "input_invalid": 0,
        "skipped_blank_rows": 0,
        "row_errors": 0,
        "valid_records": 1,
    },
    "partial-errors": {
        "status": "passed",
        "total": 4,
        "valid": 1,
        "input_invalid": 3,
        "skipped_blank_rows": 2,
        "row_errors": 3,
        "valid_records": 1,
    },
    "blocked": {
        "status": "blocked",
        "total": 1,
        "valid": 0,
        "input_invalid": 1,
        "skipped_blank_rows": 1,
        "row_errors": 1,
        "valid_records": 0,
    },
    "longest": {
        "status": "passed",
        "total": 2,
        "valid": 2,
        "input_invalid": 0,
        "skipped_blank_rows": 0,
        "row_errors": 0,
        "valid_records": 2,
    },
    "forced-failure": {
        "status": "passed",
        "total": 1,
        "valid": 1,
        "input_invalid": 0,
        "skipped_blank_rows": 0,
        "row_errors": 0,
        "valid_records": 1,
    },
}

#: five-scenarios 样例的编号 → 模拟情境；与 samples.py 中的注释一致，防止漂移。
FIVE_SCENARIO_MAP: dict[str, str] = {
    "1": "correction",
    "2": "wrong",
    "3": "insufficient",
    "6": "correct",
    "9": "forced_review",
}


def _parsed(tmp_path: Path, name: str):  # type: ignore[no-untyped-def]
    sample = get_sample(name)
    assert sample is not None, f"样例不存在：{name}"
    path = tmp_path / sample.filename
    path.write_bytes(build_workbook_bytes(sample))
    return precheck(path)


def test_sample_names_are_unique() -> None:
    names = [sample.name for sample in SAMPLES]
    assert len(names) == len(set(names))
    assert all(name == name.strip() for name in names)


def test_every_sample_declares_its_expected_outcome() -> None:
    """新增样例必须同时在 EXPECTED_OUTCOME 里声明结果，否则参数化测试会漏掉它。"""
    assert {sample.name for sample in SAMPLES} == set(EXPECTED_OUTCOME)


def test_every_sample_builds_a_readable_workbook() -> None:
    """下载到的字节必须是可被 openpyxl 打开的工作簿，表名、表头与内容一致。

    注意：全空行写入后是空的 ``<row/>``，普通读取是否把它算进 ``max_row`` 取决于
    它是否位于末行，因此这里只比较**非空行**的内容与条数，不比较行号，避免把这个
    openpyxl 差异写成脆弱的断言。空白行由预检的 ``skipped_blank_rows`` 覆盖。
    """
    for sample in SAMPLES:
        workbook = load_workbook(BytesIO(build_workbook_bytes(sample)))
        assert workbook.sheetnames == [PREFERRED_SHEET]
        sheet = workbook[PREFERRED_SHEET]
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        assert rows[0] == INPUT_COLUMNS

        def kept(row: tuple[object, ...]) -> bool:
            return any(cell is not None and str(cell).strip() for cell in row)

        written = [row for row in rows[1:] if kept(row)]
        declared = [row for row in sample.rows if kept(row)]
        assert written == declared
        assert len(written) == sample.record_count


@pytest.mark.parametrize("name", sorted(EXPECTED_OUTCOME))
def test_record_count_matches_precheck_total(tmp_path: Path, name: str) -> None:
    """``record_count`` 与界面上预检显示的 ``counts.total`` 同口径。"""
    parsed = _parsed(tmp_path, name)
    sample = get_sample(name)
    assert sample is not None
    assert parsed.report.counts.total == sample.record_count
    assert sample.record_count == EXPECTED_OUTCOME[name]["total"]


@pytest.mark.parametrize("name", sorted(EXPECTED_OUTCOME))
def test_precheck_outcome_matches_declared_description(tmp_path: Path, name: str) -> None:
    parsed = _parsed(tmp_path, name)
    report = parsed.report
    expected = EXPECTED_OUTCOME[name]

    assert str(report.status) == expected["status"]
    assert report.counts.total == expected["total"]
    assert report.counts.valid == expected["valid"]
    assert report.counts.input_invalid == expected["input_invalid"]
    assert report.counts.skipped_blank_rows == expected["skipped_blank_rows"]
    assert len(report.row_errors) == expected["row_errors"]
    assert len(parsed.valid_records()) == expected["valid_records"]


def test_blocked_sample_reports_blockers_and_no_records(tmp_path: Path) -> None:
    """blocked 样例必须真的被阻断：界面据此显示"不调用模型"。"""
    parsed = _parsed(tmp_path, "blocked")
    assert parsed.report.status == "blocked"
    assert parsed.report.blockers
    assert parsed.valid_records() == []


def test_no_refs_sample_keeps_empty_refs(tmp_path: Path) -> None:
    """资料全空的记录仍是有效记录：资料不足是要判别的结果，不是输入错误。"""
    parsed = _parsed(tmp_path, "no-refs")
    record = parsed.valid_records()[0]
    assert record.refs == {field: None for field in record.refs}
    assert record.q and record.a


def test_five_scenarios_sample_covers_every_mock_scenario(tmp_path: Path) -> None:
    """编号 1/2/3/6/9 必须恰好覆盖五类情境，界面里依次判就能看全。"""
    parsed = _parsed(tmp_path, "five-scenarios")
    mapping = {record.record_id: scenario_for(record) for record in parsed.valid_records()}
    assert mapping == FIVE_SCENARIO_MAP
    assert set(mapping.values()) == set(SCENARIOS)
    assert len(set(mapping.values())) == len(SCENARIOS)


def test_longest_sample_has_the_most_input_text(tmp_path: Path) -> None:
    """S03-06 的真实调用要求"含 1 条最长样例"，本断言守住这份样例确实最长。"""

    def text_chars(name: str) -> int:
        parsed = _parsed(tmp_path, name)
        return sum(
            len(record.q or "") + len(record.a or "") + sum(len(v or "") for v in record.refs.values())
            for record in parsed.records
        )

    longest = text_chars("longest")
    assert all(longest > text_chars(sample.name) for sample in SAMPLES if sample.name != "longest")


def test_longest_sample_has_a_multi_line_ref(tmp_path: Path) -> None:
    """最长样例的多行长资料用于上下文边界检查，不能退化成单行。"""
    parsed = _parsed(tmp_path, "longest")
    record = parsed.valid_records()[0]
    ref1 = record.refs["ref1"] or ""
    assert ref1.count("\n") >= 19


def test_get_sample_unknown_returns_none() -> None:
    assert get_sample("does-not-exist") is None
