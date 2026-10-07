"""S07-01：人工核定数据的读取与校验。

覆盖合同本身（列、标签、排除、元数据）与**一次报出全部问题**的行为；
合成样例来自 `tests/fixtures/gold_samples.py`，不含任何真实问答。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from fixtures.gold_samples import METADATA, SYNTHETIC_SOURCE_SHA256, gold_row, write_synthetic_gold
from openpyxl import Workbook

from aidhu_om_agent.evaluation.gold import (
    EXCLUDED_YES,
    GOLD_COLUMNS,
    GOLD_EXCLUDED_COLUMN,
    GOLD_EXCLUSION_REASON_COLUMN,
    GOLD_JUDGEMENT_COLUMN,
    GOLD_KEY_REF_COLUMN,
    GOLD_META_SHEET,
    GOLD_REASON_COLUMN,
    GOLD_SHEET,
    GoldReadError,
    GoldValidationError,
    compare_ids,
    read_gold,
    source_sha256_prefix,
)
from aidhu_om_agent.schemas.judgement import Label
from aidhu_om_agent.schemas.qa import A_COLUMN, ID_COLUMN, INPUT_COLUMNS, Q_COLUMN


def codes(exc: GoldValidationError) -> list[str]:
    return [problem.code for problem in exc.problems]


def write_raw(path: Path, header: list[object], rows: list[list[object]], *, meta: bool = True) -> Path:
    """直接写一个工作簿，用于构造列缺失/多余列这类样例生成器做不出的形状。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = GOLD_SHEET
    sheet.append(header)
    for row in rows:
        sheet.append(row)

    if meta:
        extra = workbook.create_sheet(title=GOLD_META_SHEET)
        extra.append(["项", "值"])
        for key, value in METADATA.items():
            extra.append([key, value])
    workbook.save(path)
    return path


# --------------------------------------------------------------------------
# 合法样例
# --------------------------------------------------------------------------


def test_synthetic_gold_reads_clean(tmp_path: Path) -> None:
    gold = read_gold(write_synthetic_gold(tmp_path / "gold.xlsx"))

    assert gold.contract_version == "1.0"
    assert gold.sheet_name == GOLD_SHEET
    assert len(gold.entries) == 12
    assert [entry.record_id for entry in gold.included] == [
        f"G-{index}" for index in range(1, 12)
    ]
    assert [entry.record_id for entry in gold.excluded] == ["G-12"]
    assert gold.label_counts == {
        Label.CORRECT.value: 6,
        Label.NO_CORRECT_REF.value: 3,
        Label.REF_OK_ANSWER_WRONG.value: 2,
    }
    assert gold.warnings == ()
    assert gold.ignored_columns == ()
    assert gold.skipped_blank_rows == 0
    assert gold.sha256_prefix == gold.sha256[:12]


def test_metadata_is_read_verbatim(tmp_path: Path) -> None:
    gold = read_gold(write_synthetic_gold(tmp_path / "gold.xlsx"))

    assert gold.metadata.data_version == "synthetic-v1"
    assert gold.metadata.annotated_by == "合成样例"
    assert gold.metadata.annotated_on == "2026-10-07"
    assert gold.metadata.source_sha256_prefix == SYNTHETIC_SOURCE_SHA256
    assert gold.metadata.as_dict()["数据版本"] == "synthetic-v1"


def test_excluded_entry_keeps_id_and_reason_only(tmp_path: Path) -> None:
    gold = read_gold(write_synthetic_gold(tmp_path / "gold.xlsx"))
    excluded = gold.excluded[0]

    assert excluded.label is None
    assert excluded.exclusion_reason == "合成理由：提问超出意图范围。"
    assert "G-12" not in gold.included_ids
    assert excluded.record_id == "G-12"
    assert excluded.source_row == 13  # 表头占第 1 行
    assert excluded.record.order_index == 11


def test_key_refs_accept_mixed_separators(tmp_path: Path) -> None:
    rows = [
        gold_row(
            "K-1",
            judgement=Label.CORRECT.value,
            key_ref="ref1、 ref2",
            refs=("合成资料一。", "合成资料二。"),
        ),
    ]
    path = write_synthetic_gold(tmp_path / "gold.xlsx", rows, metadata=METADATA)

    assert read_gold(path).entries[0].key_refs == ("ref1", "ref2")


def test_blank_rows_are_skipped(tmp_path: Path) -> None:
    rows = [
        gold_row("B-1", judgement=Label.CORRECT.value),
        {column: None for column in GOLD_COLUMNS},
        gold_row("B-2", judgement=Label.CORRECT.value),
    ]
    path = write_synthetic_gold(tmp_path / "gold.xlsx", rows, metadata=METADATA)
    gold = read_gold(path)

    assert [entry.record_id for entry in gold.included] == ["B-1", "B-2"]
    assert gold.skipped_blank_rows == 1
    assert [entry.record.order_index for entry in gold.included] == [0, 1]


def test_extra_columns_are_ignored_not_fatal(tmp_path: Path) -> None:
    """原始文件可能带「判断」这类旧列；不影响核定列，只记录被忽略的表头。"""
    rows = [[f"E-{index}", "合成提问", "合成回答", *([None] * 10),
             Label.CORRECT.value, None, None, None, None, "旧判断列"] for index in (1, 2)]
    path = write_raw(
        tmp_path / "gold.xlsx",
        [*GOLD_COLUMNS, "原判断"],
        rows,
    )
    gold = read_gold(path)

    assert gold.ignored_columns == ("原判断",)
    assert len(gold.included) == 2


def test_missing_reason_is_a_warning_not_an_error(tmp_path: Path) -> None:
    """plan/04 §1 要求保存人工理由，但缺了不影响计分，因此只报警告。"""
    rows = [gold_row("W-1", judgement=Label.CORRECT.value, reason=None)]
    gold = read_gold(write_synthetic_gold(tmp_path / "gold.xlsx", rows, metadata=METADATA))

    assert [issue.code for issue in gold.warnings] == ["GOLD_REASON_MISSING"]
    assert gold.warnings[0].source_row == 2
    assert len(gold.included) == 1


def test_sheet_without_metadata_sheet_is_rejected(tmp_path: Path) -> None:
    path = write_raw(
        tmp_path / "gold.xlsx",
        list(GOLD_COLUMNS),
        [list(gold_row("M-1", judgement=Label.CORRECT.value).values())],
        meta=False,
    )

    with pytest.raises(GoldValidationError) as info:
        read_gold(path)

    assert "GOLD_META_MISSING" in codes(info.value)


# --------------------------------------------------------------------------
# 逐项拒绝
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("rows", "expected_code"),
    [
        # 未核定：既无标签，也没标排除
        ([gold_row("X-1")], "GOLD_UNDECIDED"),
        # 近似标签不是合法标签
        ([gold_row("X-1", judgement="正确")], "GOLD_LABEL_INVALID"),
        # 缺编号
        ([gold_row(None, judgement=Label.CORRECT.value)], "GOLD_ID_MISSING"),
        # 是否排除列写了别的值
        (
            [gold_row("X-1", judgement=Label.CORRECT.value, excluded="不确定")],
            "GOLD_EXCLUDED_INVALID",
        ),
        # 排除行必须给理由
        (
            [gold_row("X-1", excluded=EXCLUDED_YES, exclusion_reason=None)],
            "GOLD_EXCLUSION_REASON_MISSING",
        ),
        # 排除行不能带标签
        (
            [gold_row("X-1", judgement=Label.CORRECT.value, excluded=EXCLUDED_YES,
                      exclusion_reason="理由")],
            "GOLD_EXCLUDED_HAS_LABEL",
        ),
        # 关键 ref 拼错
        (
            [gold_row("X-1", judgement=Label.CORRECT.value, key_ref="ref11")],
            "GOLD_KEY_REF_UNKNOWN",
        ),
        # 关键 ref 指向空的资料列
        (
            [gold_row("X-1", judgement=Label.CORRECT.value, key_ref="ref3"),
             gold_row("X-2", judgement=Label.CORRECT.value)],
            "GOLD_KEY_REF_EMPTY",
        ),
        # 没有保留输入原件
        ([gold_row("X-1", judgement=Label.CORRECT.value, q=None)], "GOLD_INPUT_FIELD_MISSING"),
        ([gold_row("X-1", judgement=Label.CORRECT.value, a=None)], "GOLD_INPUT_FIELD_MISSING"),
    ],
)
def test_row_level_rejections(tmp_path: Path, rows, expected_code: str) -> None:
    path = write_synthetic_gold(tmp_path / "gold.xlsx", rows, metadata=METADATA)

    with pytest.raises(GoldValidationError) as info:
        read_gold(path)

    assert expected_code in codes(info.value)


def test_duplicate_id_is_reported_with_first_row(tmp_path: Path) -> None:
    rows = [
        gold_row("D-1", judgement=Label.CORRECT.value),
        gold_row("D-2", judgement=Label.CORRECT.value),
        gold_row("D-1", judgement=Label.NO_CORRECT_REF.value),
    ]
    path = write_synthetic_gold(tmp_path / "gold.xlsx", rows, metadata=METADATA)

    with pytest.raises(GoldValidationError) as info:
        read_gold(path)

    duplicate = next(
        problem for problem in info.value.problems if problem.code == "GOLD_ID_DUPLICATE"
    )
    assert duplicate.source_row == 4
    assert "第 2 行" in duplicate.message


def test_all_problems_are_reported_at_once(tmp_path: Path) -> None:
    """用户要一次改完 100 行：三条坏行必须一次全报，而不是只报第一条。"""
    rows = [
        gold_row("P-1"),
        gold_row("P-2", judgement="回答对"),
        gold_row("P-3", excluded=EXCLUDED_YES),
    ]
    path = write_synthetic_gold(tmp_path / "gold.xlsx", rows, metadata=METADATA)

    with pytest.raises(GoldValidationError) as info:
        read_gold(path)

    assert codes(info.value) == ["GOLD_UNDECIDED", "GOLD_LABEL_INVALID", "GOLD_EXCLUSION_REASON_MISSING"]
    assert [problem.source_row for problem in info.value.problems] == [2, 3, 4]
    assert str(info.value).startswith("核定文件有 3 处问题：")


def test_missing_columns_list_every_missing_one(tmp_path: Path) -> None:
    header = [ID_COLUMN, Q_COLUMN, A_COLUMN, *INPUT_COLUMNS[3:]]
    path = write_raw(
        tmp_path / "gold.xlsx",
        header,
        [["C-1", "合成提问", "合成回答", *([None] * 10)]],
    )

    with pytest.raises(GoldValidationError) as info:
        read_gold(path)

    problem = next(
        item for item in info.value.problems if item.code == "GOLD_MISSING_COLUMNS"
    )
    for column in (
        GOLD_JUDGEMENT_COLUMN,
        GOLD_REASON_COLUMN,
        GOLD_KEY_REF_COLUMN,
        GOLD_EXCLUDED_COLUMN,
        GOLD_EXCLUSION_REASON_COLUMN,
    ):
        assert column in problem.message


def test_empty_sheet_is_rejected(tmp_path: Path) -> None:
    path = write_raw(tmp_path / "gold.xlsx", list(GOLD_COLUMNS), [])

    with pytest.raises(GoldValidationError) as info:
        read_gold(path)

    assert "GOLD_NO_ROWS" in codes(info.value)


@pytest.mark.parametrize(
    ("metadata", "expected_code"),
    [
        ({**METADATA, "核定日期": "2026/10/07"}, "GOLD_META_INVALID_DATE"),
        ({**METADATA, "原始文件sha256前12位": "0123456789ab-cd"}, "GOLD_META_INVALID_SHA256"),
        ({**METADATA, "核定人": None}, "GOLD_META_INCOMPLETE"),
        ({**METADATA, "原始文件sha256前12位": "0123456789AB"}, None),  # 大写可接受，读回统一小写
    ],
)
def test_metadata_format(tmp_path: Path, metadata, expected_code: str | None) -> None:
    rows = [gold_row("M-1", judgement=Label.CORRECT.value)]
    path = write_synthetic_gold(tmp_path / "gold.xlsx", rows, metadata=metadata)

    if expected_code is None:
        assert read_gold(path).metadata.source_sha256_prefix == SYNTHETIC_SOURCE_SHA256
        return

    with pytest.raises(GoldValidationError) as info:
        read_gold(path)

    assert expected_code in codes(info.value)


# --------------------------------------------------------------------------
# 读不了 vs 不合合同
# --------------------------------------------------------------------------


def test_missing_file_is_a_read_error(tmp_path: Path) -> None:
    with pytest.raises(GoldReadError):
        read_gold(tmp_path / "没有这个文件.xlsx")


def test_unreadable_file_is_a_read_error(tmp_path: Path) -> None:
    path = tmp_path / "not-a-workbook.xlsx"
    path.write_text("这不是 xlsx", encoding="utf-8")

    with pytest.raises(GoldReadError):
        read_gold(path)


def test_ambiguous_main_sheet_is_a_read_error(tmp_path: Path) -> None:
    """元数据表之外还有两个可见表、且没有 QA_REF 时无法确定主表。"""
    workbook = Workbook()
    workbook.active.title = "表一"
    workbook.active.append(list(GOLD_COLUMNS))
    workbook.create_sheet(title="表二").append(list(GOLD_COLUMNS))
    meta = workbook.create_sheet(title=GOLD_META_SHEET)
    meta.append(["项", "值"])
    for key, value in METADATA.items():
        meta.append([key, value])
    path = tmp_path / "gold.xlsx"
    workbook.save(path)

    with pytest.raises(GoldReadError) as info:
        read_gold(path)

    assert "无法确定核定工作表" in str(info.value)


# --------------------------------------------------------------------------
# 编号比对与摘要工具
# --------------------------------------------------------------------------


def test_compare_ids_returns_both_differences() -> None:
    missing, extra = compare_ids(["a", "b", "c"], ["b", "c", "d"])

    assert missing == ("a",)
    assert extra == ("d",)


def test_compare_ids_is_empty_when_identical() -> None:
    assert compare_ids(["a", "b"], ["b", "a"]) == ((), ())


def test_source_sha256_prefix_normalises_case_and_length() -> None:
    digest = "AB" * 32

    assert source_sha256_prefix(digest) == "ab" * 6
    assert len(source_sha256_prefix(digest)) == 12
