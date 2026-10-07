"""S05-01 导出字段与编码合同的单测。

覆盖四张表的列序、人读序号、长文本截断、公式注入防护、控制字符清理、空值留空、
下载文件名规则，以及写出层「单元格里不允许出现公式」的硬规则。

**不读数据库、不生成业务数据**：只验证编码后的字符串与写出的文件。全程零真实调用。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from aidhu_om_agent.excel.writer import ExcelWriteError, SheetData, write_workbook
from aidhu_om_agent.schemas.export import (
    ARTIFACT_EXCEL,
    ARTIFACT_JSONL,
    CLASSIFICATION_COLUMNS,
    EXCEL_CELL_MAX_CHARS,
    FAILURE_COLUMNS,
    REVIEW_COLUMNS,
    REVIEW_HUMAN_COLUMNS,
    SHEET_CLASSIFICATION,
    SHEET_FAILURES,
    SHEET_ORDER,
    SHEET_REVIEW,
    SHEET_SUMMARY,
    SUMMARY_COLUMNS,
    SUMMARY_FIELDS,
    TEXT_GUARD_PREFIX,
    TRUNCATION_MARKER,
    cell_location,
    download_name,
    excel_cell_text,
    format_evidence,
    ordinal,
    record_status_text,
    retryable_text,
    review_required_text,
    run_status_text,
    sanitized_cell_location,
    stage_text,
    strip_illegal_cell_chars,
    truncate_for_excel,
)
from aidhu_om_agent.schemas.qa import REF_FIELDS


# ------------------------------------------------------------------ 列序合同


def test_sheet_order_and_names_are_fixed() -> None:
    assert SHEET_ORDER == (
        "分类结果",
        "复核清单",
        "失败清单",
        "运行概况",
    )


def test_classification_columns_match_plan_05_section_6() -> None:
    """plan/05 §6「分类结果」：编号、原预测、理由、证据、复核标记和处理状态。"""
    assert CLASSIFICATION_COLUMNS == (
        "序号",
        "原始编号",
        "来源行",
        "原预测",
        "理由",
        "证据",
        "需复核",
        "处理状态",
    )


def test_review_columns_carry_all_ten_refs_and_two_blank_human_columns() -> None:
    assert REVIEW_COLUMNS[:5] == ("序号", "原始编号", "来源行", "q", "a")
    assert REVIEW_COLUMNS[5 : 5 + len(REF_FIELDS)] == REF_FIELDS
    assert len(REF_FIELDS) == 10
    # 两个人工填写列在最后，导出时留空、不回读覆盖原预测。
    assert REVIEW_COLUMNS[-2:] == REVIEW_HUMAN_COLUMNS == ("人工判断", "复核说明")


def test_failure_columns_match_plan_05_section_6() -> None:
    assert FAILURE_COLUMNS == (
        "序号",
        "原始编号",
        "来源行",
        "失败阶段",
        "错误码",
        "错误信息",
        "尝试次数",
        "是否可重试",
    )


def test_summary_is_a_two_column_key_value_sheet() -> None:
    assert SUMMARY_COLUMNS == ("项目", "值")
    # 「未处理」必须独立成行：它是「这份导出不是完整批次」的唯一线索。
    assert "未处理" in SUMMARY_FIELDS
    assert "被截断单元格数" in SUMMARY_FIELDS
    assert len(set(SUMMARY_FIELDS)) == len(SUMMARY_FIELDS)


# ------------------------------------------------------------------ 人读序号


def test_ordinal_is_one_based_and_leaves_order_index_untouched() -> None:
    assert ordinal(0) == 1
    assert ordinal(7) == 8


# ------------------------------------------------------------------ 长文本


def test_text_at_the_limit_is_not_truncated() -> None:
    text = "字" * EXCEL_CELL_MAX_CHARS
    kept, truncated = truncate_for_excel(text)
    assert truncated is False
    assert kept == text
    assert excel_cell_text(text).truncated is False


def test_text_over_the_limit_is_truncated_with_a_visible_marker() -> None:
    text = "字" * (EXCEL_CELL_MAX_CHARS + 1)
    cell = excel_cell_text(text)
    assert cell.truncated is True
    assert len(cell.text) <= EXCEL_CELL_MAX_CHARS
    assert cell.text.endswith(TRUNCATION_MARKER)
    # 头部原样保留：截断只切尾部，不改写前面的内容。
    head = EXCEL_CELL_MAX_CHARS - len(TRUNCATION_MARKER)
    assert cell.text[:head] == text[:head]


def test_cell_location_names_sheet_column_row_and_original_length() -> None:
    assert cell_location("分类结果", "理由", 3, 41234) == (
        "分类结果!理由 第 3 行（原 41234 字符）"
    )


# ------------------------------------------------------------ 公式注入防护


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_formula_like_text_is_prefixed_so_excel_treats_it_as_text(prefix: str) -> None:
    cell = excel_cell_text(f"{prefix}1+1")
    assert cell.guarded is True
    assert cell.text == f"{TEXT_GUARD_PREFIX}{prefix}1+1"


def test_ordinary_text_is_not_prefixed() -> None:
    cell = excel_cell_text("正常的资料文本")
    assert cell.guarded is False
    assert cell.text == "正常的资料文本"


def test_guard_and_truncation_together_stay_within_the_cell_limit() -> None:
    """先加前缀再截断：否则 32767 字符的前缀结果会变成 32768，写入再次越界。"""
    text = "=" + "x" * (EXCEL_CELL_MAX_CHARS * 2)
    cell = excel_cell_text(text)
    assert cell.guarded is True
    assert cell.truncated is True
    assert len(cell.text) <= EXCEL_CELL_MAX_CHARS
    assert cell.text.startswith(TEXT_GUARD_PREFIX)
    assert cell.text.endswith(TRUNCATION_MARKER)


def test_control_characters_before_a_formula_char_still_trigger_the_guard() -> None:
    """控制字符排在开头时不能顶掉防护判断，否则 `\\x0b=` 会被漏掉。"""
    cell = excel_cell_text("\x0b=1+1")
    assert cell.sanitized == 1
    assert cell.guarded is True
    assert cell.text == f"{TEXT_GUARD_PREFIX}=1+1"


# ------------------------------------------------------------ 控制字符清理


def test_illegal_control_characters_are_removed_with_a_count() -> None:
    text, removed = strip_illegal_cell_chars("a\x00b\x0bc\x1fd")
    assert text == "abcd"
    assert removed == 3


def test_tab_and_newline_are_legal_and_kept() -> None:
    text = "第一行\n第二行\t带制表符"
    assert strip_illegal_cell_chars(text) == (text, 0)
    assert excel_cell_text(text).text == text


def test_sanitized_location_reports_the_removed_count() -> None:
    assert sanitized_cell_location("复核清单", "ref3", 7, 2) == (
        "复核清单!ref3 第 7 行（清理 2 个控制字符）"
    )


# ------------------------------------------------------------ 空值与展示映射


def test_missing_values_are_left_blank_not_filled_with_placeholders() -> None:
    for value in (None, ""):
        cell = excel_cell_text(value)
        assert cell.text == ""
        assert cell.rewritten is False


def test_status_and_flag_columns_render_chinese_or_blank() -> None:
    assert record_status_text("completed") == "已分类"
    assert record_status_text("failed") == "技术失败"
    assert record_status_text("input_invalid") == "输入失败"
    assert record_status_text("pending") == "未处理"
    assert review_required_text(1) == "是"
    assert review_required_text(0) == "否"
    assert review_required_text(None) == ""
    assert retryable_text(True) == "是"
    assert retryable_text(False) == "否"
    assert retryable_text(None) == ""
    assert stage_text("stage1") == "阶段一"
    assert stage_text("stage2") == "阶段二"
    assert stage_text(None) == ""
    # 状态写成「中文（原值）」，让概况表能与库里的枚举逐字核对。
    assert run_status_text("partial_failed") == "部分失败（partial_failed）"


def test_evidence_is_rendered_as_ref_id_plus_quote_per_line() -> None:
    assert format_evidence(
        [{"ref_id": "ref3", "quote": "原文一"}, {"ref_id": "ref5", "quote": "原文二"}]
    ) == "ref3：原文一\nref5：原文二"
    assert format_evidence([]) == ""
    assert format_evidence(None) == ""


# ------------------------------------------------------------------ 文件名


def test_download_names_use_readable_stem_plus_short_ids() -> None:
    excel = download_name(ARTIFACT_EXCEL, run_id="a" * 32, export_id="b" * 32)
    jsonl = download_name(ARTIFACT_JSONL, run_id="a" * 32, export_id="b" * 32)
    assert excel == f"分类结果-{'a' * 8}-{'b' * 8}.xlsx"
    assert jsonl == f"两阶段明细-{'a' * 8}-{'b' * 8}.jsonl"


def test_download_name_rejects_unknown_kind() -> None:
    with pytest.raises(ValueError):
        download_name("csv", run_id="a", export_id="b")


# ------------------------------------------------------------------ 写出层


def test_writer_refuses_an_unguarded_formula_cell(tmp_path: Path) -> None:
    """编码漏做的字符串必须当场失败，不能生成一份打开就求值的文件。"""
    sheets = [SheetData("表", ("值",), (("=1+1",),))]
    with pytest.raises(ExcelWriteError, match="文本前缀防护"):
        write_workbook(tmp_path / "bad.xlsx", sheets)


def test_writer_rejects_ragged_rows(tmp_path: Path) -> None:
    sheets = [SheetData("表", ("甲", "乙"), (("只有一列",),))]
    with pytest.raises(ExcelWriteError, match="表头是 2 列"):
        write_workbook(tmp_path / "bad.xlsx", sheets)


def test_written_workbook_round_trips_four_sheets_as_text(tmp_path: Path) -> None:
    target = tmp_path / "out.xlsx"
    guarded = excel_cell_text("=1+1").text
    sheets = (
        SheetData(SHEET_CLASSIFICATION, ("序号", "值"), ((1, guarded),)),
        SheetData(SHEET_REVIEW, ("序号", "值"), ()),
        SheetData(SHEET_FAILURES, ("序号", "值"), ()),
        SheetData(SHEET_SUMMARY, ("项目", "值"), (("导出合同版本", "1.0"),)),
    )
    write_workbook(target, sheets)

    workbook = load_workbook(target)
    assert workbook.sheetnames == list(SHEET_ORDER)
    sheet = workbook[SHEET_CLASSIFICATION]
    assert sheet["A1"].value == "序号"
    assert sheet["A2"].value == 1
    # 写进去的是**文本**：前缀保留、不会被求值成 2。
    assert sheet["B2"].value == guarded
    assert sheet["B2"].data_type == "s"
    assert workbook[SHEET_SUMMARY]["B2"].value == "1.0"
    assert workbook[SHEET_REVIEW].max_row == 1  # 只有表头
